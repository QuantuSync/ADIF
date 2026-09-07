"""Bloque 1, sesión del Excel de ejecución SAP (2026-09-07): estado de
contrato exportado directamente de SAP por ADIF (`EXPEDIENTES_EJECUCION_SAP`,
columnas "Expediente ADIF" / "Título del expediente" / "Descripción del
estado") -- un dato de negocio que no existe en ningún documento de la
Plataforma ni en la sindicación, así que este sistema no puede deducirlo
solo. Se trata como fuente de entrada permanente, igual que el Excel de
códigos (`app.extraccion.cruce_codigos`): una ruta configurada
(`ESTADO_SAP_PATH`), montada por bind-mount, que se puede volver a cargar
cuantas veces haga falta.

Carga por `codigo_expediente` exacto (misma normalización que
`cruce_codigos.normalizar_codigo_expediente`), nunca por título -- dos
expedientes de suministros distintos pueden compartir casi el mismo título
("SUMINISTRO... LOTE 3" / "...LOTE 4"; el propio Excel de ejemplo de esta
sesión trae varios), así que cruzar por texto inventaría relaciones falsas
(CONTEXTO.md sección 7: "el cruce es por clave exacta, no por similitud de
nombre").

Repetible por diseño (encargo de esta sesión: "si mañana pasan una
exportación nueva, se actualiza sin duplicar nada"): upsert por
`codigo_expediente`, mismo mecanismo idempotente que
`app.sindicacion.descubrimiento.descubrir_novedades` y `POST /expedientes`
(CONTEXTO.md sección 9.9). Un expediente del Excel de SAP que el sistema
todavía no conocía se da de alta con el estado de PROCESAMIENTO de siempre
(`pendiente`, valor por defecto de `Expediente.estado`) -- este módulo solo
escribe el estado de CONTRATO; no descarga ni encola nada él mismo, igual
que `descubrir_novedades` tampoco descarga (bloque 3 se encarga de encolar
lo que falte)."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.extraccion.cruce_codigos import normalizar_codigo_expediente
from app.models import Expediente

COLUMNA_EXPEDIENTE = "Expediente ADIF"
COLUMNA_TITULO = "Título del expediente"
COLUMNA_ESTADO = "Descripción del estado"


class EstadoSapPathInvalida(RuntimeError):
    """Mismo criterio que `CodigosProyectoPathInvalida`
    (`app.extraccion.cruce_codigos`): fallar de forma visible al validar la
    ruta en vez de dejar la carga rota para el primer intento real."""


def validar_ruta_estado_sap(ruta: Optional[str]) -> None:
    """Se llama una vez al arrancar la API (`app.main`) -- sin ruta
    configurada no hay nada que validar, la carga manual simplemente no
    tiene nada que ofrecer (`ResumenCargaEstadoSap` con `configurado=False`)."""
    if not ruta:
        return
    camino = Path(ruta)
    if not camino.is_file():
        raise EstadoSapPathInvalida(
            f"ESTADO_SAP_PATH={ruta!r} no es un fichero (¿bind-mount con el origen ausente?)"
        )
    try:
        libro = openpyxl.load_workbook(camino, read_only=True)
        libro.close()
    except Exception as exc:
        raise EstadoSapPathInvalida(
            f"ESTADO_SAP_PATH={ruta!r} no se puede abrir como .xlsx: {exc}"
        ) from exc


@dataclass
class ResumenCargaEstadoSap:
    configurado: bool
    filas_leidas: int = 0
    filas_sin_codigo: int = 0
    expedientes_nuevos: int = 0
    expedientes_actualizados: int = 0
    expedientes_sin_cambios: int = 0

    def to_dict(self) -> dict:
        return {
            "configurado": self.configurado,
            "filas_leidas": self.filas_leidas,
            "filas_sin_codigo": self.filas_sin_codigo,
            "expedientes_nuevos": self.expedientes_nuevos,
            "expedientes_actualizados": self.expedientes_actualizados,
            "expedientes_sin_cambios": self.expedientes_sin_cambios,
        }


def _texto(valor: object) -> Optional[str]:
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def cargar_estado_sap(db: Session, ruta_excel: Optional[str]) -> ResumenCargaEstadoSap:
    """Lee todas las hojas del Excel (mismo motivo que
    `cruce_codigos._cargar_indice`: no asumir que el dato vive siempre en la
    primera) y hace upsert de `estado_contrato_sap` por `codigo_expediente`
    normalizado. `nombre_proyecto` solo se rellena si estaba vacío -- nunca
    pisa el nombre ya extraído del "Objeto del Contrato" de un PDF real
    (CONTEXTO.md sección 7), que es la fuente autorizada; el título de SAP es
    solo un mejor-que-nada para un expediente recién descubierto que aún no
    se ha procesado."""
    resumen = ResumenCargaEstadoSap(configurado=bool(ruta_excel))
    if not ruta_excel:
        return resumen

    libro = openpyxl.load_workbook(ruta_excel, read_only=True, data_only=True)
    ahora = datetime.now(timezone.utc)
    try:
        for hoja in libro.worksheets:
            filas_iter = hoja.iter_rows(values_only=True)
            primera = next(filas_iter, None)
            if primera is None:
                continue  # hoja vacía, sin cabecera siquiera
            cabecera = [str(c).strip() if c is not None else "" for c in primera]
            if COLUMNA_EXPEDIENTE not in cabecera:
                continue  # hoja sin esta columna: no es la hoja de expedientes en ejecución
            indice_codigo = cabecera.index(COLUMNA_EXPEDIENTE)
            indice_titulo = cabecera.index(COLUMNA_TITULO) if COLUMNA_TITULO in cabecera else None
            indice_estado = cabecera.index(COLUMNA_ESTADO) if COLUMNA_ESTADO in cabecera else None

            for fila in filas_iter:
                resumen.filas_leidas += 1
                codigo = normalizar_codigo_expediente(fila[indice_codigo] if indice_codigo < len(fila) else None)
                if not codigo:
                    resumen.filas_sin_codigo += 1
                    continue
                titulo = _texto(fila[indice_titulo]) if indice_titulo is not None and indice_titulo < len(fila) else None
                estado_sap = _texto(fila[indice_estado]) if indice_estado is not None and indice_estado < len(fila) else None

                expediente = db.execute(
                    select(Expediente).where(Expediente.codigo_expediente == codigo)
                ).scalar_one_or_none()
                if expediente is None:
                    expediente = Expediente(codigo_expediente=codigo, nombre_proyecto=titulo)
                    db.add(expediente)
                    resumen.expedientes_nuevos += 1
                else:
                    if expediente.nombre_proyecto is None and titulo:
                        expediente.nombre_proyecto = titulo
                    if expediente.estado_contrato_sap == estado_sap:
                        resumen.expedientes_sin_cambios += 1
                    else:
                        resumen.expedientes_actualizados += 1
                expediente.estado_contrato_sap = estado_sap
                expediente.estado_contrato_sap_actualizado_en = ahora
            db.commit()
    finally:
        libro.close()
    return resumen

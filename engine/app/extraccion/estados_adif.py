"""Listado de estados de contratación que ADIF envió el 18/09/2026 (bloque 1
del encargo de esa sesión).

**Procedencia, escrita aquí para que no haya que preguntarla.** Nos lo envió
ADIF el 18/09/2026. Su origen es una transacción de SAP ejecutada **por
ellos**: nosotros no lo hemos sacado de ningún sistema. En su día se nos
indicó no utilizar volcados de SAP, así que se preguntó expresamente en el
grupo si podíamos usar este, y **nos autorizaron a usarlo**.

**Qué trae.** Una sola hoja, cuatro columnas -- "Título del expediente",
"Expediente ADIF", "Fecha de creación" y "Descripción del estado" -- y 358
filas, todas del departamento 28510, sin ningún código repetido. **No trae
presupuesto de licitación ni órgano de contratación**: ninguna de las dos
cosas se puede rellenar desde aquí.

**Qué NO puede hacer este módulo, y es la razón de que no reutilice
`app.extraccion.estado_sap`:**

1. **No decide si un expediente está publicado.** Esa lógica se apoya solo en
   la Plataforma (`app.conciliacion.consta_publicado`: sindicación, búsqueda
   directa o documento descargado). Este listado es fuente de contraste y de
   relleno de columnas, nada más -- criterio explícito del cliente.
2. **No da de alta ningún expediente.** `estado_sap.cargar_estado_sap` sí lo
   hace (creó 327 en su día, y está bien que lo haga: ahí el Excel de SAP era
   la única fuente de la que el sistema podía saber que esos expedientes
   existían). Aquí, dar de alta un expediente convertiría la pregunta "¿cuáles
   de los suyos no tenemos?" en una que ya no se puede contestar, porque el
   propio acto de leer el fichero la respondería que ninguno. Los códigos que
   no casan se cuentan y se devuelven, no se crean.
3. **No escribe en `estado_contrato_sap`.** Las dos fuentes son volcados de
   SAP y su vocabulario de estados coincide, pero son volcados distintos con
   selecciones distintas de expedientes (medido: 212 códigos en común, 155
   solo en el anterior, 146 solo en este; en los 212 comunes el estado
   coincide en los 212). Compartir campo dejaría el valor sin procedencia.

Cruce por `codigo_expediente` exacto (misma normalización que
`cruce_codigos.normalizar_codigo_expediente`), nunca por título: dos
suministros distintos pueden compartir casi el mismo texto (CONTEXTO.md
sección 7, "el cruce es por clave exacta, no por similitud de nombre").

Repetible: upsert por `codigo_expediente`, igual que el resto de fuentes de
entrada permanentes. Sustituir el fichero y volver a llamar a
`POST /mantenimiento/estados-adif/cargar`.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Optional

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.extraccion.cruce_codigos import normalizar_codigo_expediente
from app.extraccion.normalizacion import parsear_importe_es
from app.extraccion.texto import normalizar
from app.models import Expediente

COLUMNA_EXPEDIENTE = "Expediente ADIF"
COLUMNA_TITULO = "Título del expediente"
COLUMNA_FECHA = "Fecha de creación"
COLUMNA_ESTADO = "Descripción del estado"

# Bloque 2, sesión 2026-09-19 (sexta parte): la columna que ADIF tiene que
# añadir a este mismo listado. **El fichero de hoy no la trae** (se comprobó
# columna a columna, CONTEXTO.md sección 7), así que su nombre exacto no se
# puede saber: se reconoce por contenido del encabezado, normalizado, en vez
# de por igualdad. "Presupuesto", "Presupuesto de licitación", "Presupuesto
# base de licitación", "Importe de licitación" y "PBL" caen todas aquí. Si el
# listado llega sin ninguna, el campo se queda vacío y no pasa nada más: la
# carga sigue funcionando igual que hasta hoy.
_ENCABEZADOS_PRESUPUESTO = ("presupuesto", "importe de licitacion", "pbl")
# Deliberadamente fuera: un importe de adjudicación no es el presupuesto de
# licitación, y confundirlos haría que la comparación de la hoja
# "Presupuestos ADIF" contrastara dos cosas distintas.
_ENCABEZADOS_PRESUPUESTO_EXCLUIDOS = ("adjudicacion", "adjudicado", "iva")

# Cuántos códigos sin expediente en el sistema se devuelven en el resumen.
# La lista completa es el bloque 2 del encargo y se mide aparte; aquí es una
# muestra para que el resultado de la carga sea legible.
_MAX_CODIGOS_DEVUELTOS = 50


class EstadosAdifPathInvalida(RuntimeError):
    """Mismo criterio que `CodigosProyectoPathInvalida` y
    `EstadoSapPathInvalida`: fallar de forma visible al validar la ruta al
    arrancar, en vez de dejar la carga rota para el primer intento real."""


def validar_ruta_estados_adif(ruta: Optional[str]) -> None:
    if not ruta:
        return
    camino = Path(ruta)
    if not camino.is_file():
        raise EstadosAdifPathInvalida(
            f"ESTADOS_ADIF_PATH={ruta!r} no es un fichero (¿bind-mount con el origen ausente?)"
        )
    try:
        libro = openpyxl.load_workbook(camino, read_only=True)
        libro.close()
    except Exception as exc:
        raise EstadosAdifPathInvalida(
            f"ESTADOS_ADIF_PATH={ruta!r} no se puede abrir como .xlsx: {exc}"
        ) from exc


@dataclass
class ResumenCargaEstadosAdif:
    configurado: bool
    filas_leidas: int = 0
    filas_sin_codigo: int = 0
    codigos_distintos: int = 0
    expedientes_actualizados: int = 0
    expedientes_sin_cambios: int = 0
    # Códigos del listado de ADIF que no existen en este sistema. **No se dan
    # de alta** (ver punto 2 del docstring del módulo): se cuentan y se
    # devuelven para que quien lo pidió decida.
    sin_expediente_en_el_sistema: int = 0
    # Bloque 2, sesión 2026-09-19 (sexta parte): cuántas filas traían
    # presupuesto de licitación, y si el listado llegó con esa columna.
    trae_presupuesto: bool = False
    presupuestos_leidos: int = 0
    codigos_sin_expediente: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "configurado": self.configurado,
            "filas_leidas": self.filas_leidas,
            "filas_sin_codigo": self.filas_sin_codigo,
            "codigos_distintos": self.codigos_distintos,
            "expedientes_actualizados": self.expedientes_actualizados,
            "expedientes_sin_cambios": self.expedientes_sin_cambios,
            "sin_expediente_en_el_sistema": self.sin_expediente_en_el_sistema,
            "trae_presupuesto": self.trae_presupuesto,
            "presupuestos_leidos": self.presupuestos_leidos,
            "codigos_sin_expediente": self.codigos_sin_expediente,
        }


def _texto(valor: object) -> Optional[str]:
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def _indice_presupuesto(cabecera: list[str]) -> Optional[int]:
    """La columna de presupuesto de licitación, si el listado la trae. Con dos
    o más candidatas no se adivina: se devuelve `None` y el campo se queda
    vacío, que es lo mismo que pasa hoy con el fichero que no la trae."""
    candidatas = [
        indice
        for indice, texto in enumerate(normalizar(c) for c in cabecera)
        if any(p in texto for p in _ENCABEZADOS_PRESUPUESTO)
        and not any(p in texto for p in _ENCABEZADOS_PRESUPUESTO_EXCLUIDOS)
    ]
    return candidatas[0] if len(candidatas) == 1 else None


def _importe(valor: object) -> Optional[Decimal]:
    """El presupuesto, venga como número de Excel o como texto en formato
    español ("1.234.567,89 €"). Lo que no se pueda interpretar se deja vacío:
    nunca se guarda una cifra a medias."""
    if valor is None:
        return None
    if isinstance(valor, (int, float, Decimal)):
        try:
            return Decimal(str(valor))
        except (ValueError, ArithmeticError):
            return None
    try:
        return parsear_importe_es(str(valor))
    except (ValueError, ArithmeticError):
        return None


def _fecha(valor: object) -> Optional[datetime]:
    """La columna "Fecha de creación" llega como `datetime` de openpyxl. Sin
    zona horaria en el fichero: se marca UTC para poder guardarla en una
    columna con zona, nunca se desplaza."""
    if isinstance(valor, datetime):
        return valor if valor.tzinfo else valor.replace(tzinfo=timezone.utc)
    return None


def _mismo_instante(a: Optional[datetime], b: Optional[datetime]) -> bool:
    """Compara dos fechas sin que la zona horaria decida por sí sola. Lo que
    se guarda viene sin zona en el fichero y se marca UTC al leerlo; lo que
    devuelve la base de datos puede venir con zona o sin ella según el motor.
    Comparar los objetos tal cual haría que una recarga idéntica se contara
    como un cambio."""
    if a is None or b is None:
        return a is None and b is None
    izquierda = a.replace(tzinfo=None) if a.tzinfo is None else a.astimezone(timezone.utc).replace(tzinfo=None)
    derecha = b.replace(tzinfo=None) if b.tzinfo is None else b.astimezone(timezone.utc).replace(tzinfo=None)
    return izquierda == derecha


def cargar_estados_adif(db: Session, ruta_excel: Optional[str]) -> ResumenCargaEstadosAdif:
    """Lee todas las hojas del fichero (mismo motivo que
    `cruce_codigos._cargar_indice`: no asumir que el dato vive siempre en la
    primera) y hace upsert de `estado_adif` sobre los expedientes que YA
    existen. El título del listado **no pisa** `nombre_proyecto`: ese sale del
    "Objeto del Contrato" de un documento real, que es la fuente autorizada
    (CONTEXTO.md sección 7); solo se usa si estaba vacío."""
    resumen = ResumenCargaEstadosAdif(configurado=bool(ruta_excel))
    if not ruta_excel:
        return resumen

    libro = openpyxl.load_workbook(ruta_excel, read_only=True, data_only=True)
    ahora = datetime.now(timezone.utc)
    vistos: set[str] = set()
    try:
        for hoja in libro.worksheets:
            filas_iter = hoja.iter_rows(values_only=True)
            primera = next(filas_iter, None)
            if primera is None:
                continue
            cabecera = [str(c).strip() if c is not None else "" for c in primera]
            if COLUMNA_EXPEDIENTE not in cabecera or COLUMNA_ESTADO not in cabecera:
                continue  # no es la hoja de estados
            i_codigo = cabecera.index(COLUMNA_EXPEDIENTE)
            i_estado = cabecera.index(COLUMNA_ESTADO)
            i_titulo = cabecera.index(COLUMNA_TITULO) if COLUMNA_TITULO in cabecera else None
            i_fecha = cabecera.index(COLUMNA_FECHA) if COLUMNA_FECHA in cabecera else None
            i_presupuesto = _indice_presupuesto(cabecera)
            resumen.trae_presupuesto = resumen.trae_presupuesto or i_presupuesto is not None

            def _celda(fila, indice):
                if indice is None or indice >= len(fila):
                    return None
                return fila[indice]

            for fila in filas_iter:
                resumen.filas_leidas += 1
                codigo = normalizar_codigo_expediente(_celda(fila, i_codigo))
                if not codigo:
                    resumen.filas_sin_codigo += 1
                    continue
                vistos.add(codigo)
                estado = _texto(_celda(fila, i_estado))
                titulo = _texto(_celda(fila, i_titulo))
                creado_en = _fecha(_celda(fila, i_fecha))
                presupuesto = _importe(_celda(fila, i_presupuesto))

                expediente = db.execute(
                    select(Expediente).where(Expediente.codigo_expediente == codigo)
                ).scalar_one_or_none()
                if expediente is None:
                    resumen.sin_expediente_en_el_sistema += 1
                    if len(resumen.codigos_sin_expediente) < _MAX_CODIGOS_DEVUELTOS:
                        resumen.codigos_sin_expediente.append(codigo)
                    continue

                if expediente.nombre_proyecto is None and titulo:
                    expediente.nombre_proyecto = titulo
                if expediente.estado_adif == estado and _mismo_instante(
                    expediente.estado_adif_creado_en, creado_en
                ):
                    resumen.expedientes_sin_cambios += 1
                else:
                    resumen.expedientes_actualizados += 1
                # El presupuesto del listado nunca pisa `importe_licitacion`
                # (que sale de los documentos publicados): campo propio, y
                # `None` no borra lo que ya hubiera de una carga anterior.
                if presupuesto is not None:
                    expediente.presupuesto_licitacion_adif = presupuesto
                    resumen.presupuestos_leidos += 1
                expediente.estado_adif = estado
                expediente.estado_adif_creado_en = creado_en
                expediente.estado_adif_actualizado_en = ahora
            db.commit()
    finally:
        libro.close()
    resumen.codigos_distintos = len(vistos)
    return resumen

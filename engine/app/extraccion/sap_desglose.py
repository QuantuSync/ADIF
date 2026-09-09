"""Desglose de SAP con las matrículas concretas de cada contrato (bloque 6,
cambios del cliente tras revisar el catálogo): ADIF ha facilitado un
Excel exportado directamente de SAP -- expediente, documento de compras,
posición, material (matrícula de 9 dígitos), texto breve, cantidad
prevista, precio neto pedido y unidad de medida base -- justo el dato que
los pliegos no traen (CONTEXTO.md sección 2: "matrícula... presente solo
en ~66% de las líneas").

Mismo mecanismo que `app.extraccion.estado_sap`/`app.extraccion.cruce_codigos`:
fuente de entrada permanente (`SAP_DESGLOSE_PATH`), montada por bind-mount,
repetible sin duplicar. A diferencia de esas dos, el cruce no es 1:1 por
expediente sino 1:N (varias líneas de compra reales por expediente) --
`sap_desglose_lineas` (migración 0024) guarda cada línea tal cual, y
`app.catalogo`/análisis posteriores deciden qué hacer con ellas; este
módulo solo carga, nunca escribe en `lineas_catalogo`."""
from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Optional

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.extraccion.cruce_codigos import normalizar_codigo_expediente
from app.models import SapDesgloseLinea

COLUMNA_EXPEDIENTE = "expediente"
COLUMNA_DOCUMENTO_COMPRAS = "Documento compras"
COLUMNA_POSICION = "Posición"
COLUMNA_MATERIAL = "Material"
COLUMNA_TEXTO_BREVE = "Texto breve"
COLUMNA_CANTIDAD_PREVISTA = "Cantidad prevista"
COLUMNA_PRECIO_NETO = "Precio neto pedido"
COLUMNA_UNIDAD_MEDIDA = "Unidad medida base"

_COLUMNAS_REQUERIDAS = (COLUMNA_EXPEDIENTE, COLUMNA_DOCUMENTO_COMPRAS, COLUMNA_POSICION)


class SapDesglosePathInvalida(RuntimeError):
    """Mismo criterio que `EstadoSapPathInvalida`/`CodigosProyectoPathInvalida`:
    fallar de forma visible al validar la ruta en vez de dejar la carga rota
    para el primer intento real."""


def validar_ruta_sap_desglose(ruta: Optional[str]) -> None:
    if not ruta:
        return
    camino = Path(ruta)
    if not camino.is_file():
        raise SapDesglosePathInvalida(
            f"SAP_DESGLOSE_PATH={ruta!r} no es un fichero (¿bind-mount con el origen ausente?)"
        )
    try:
        libro = openpyxl.load_workbook(camino, read_only=True)
        libro.close()
    except Exception as exc:
        raise SapDesglosePathInvalida(
            f"SAP_DESGLOSE_PATH={ruta!r} no se puede abrir como .xlsx: {exc}"
        ) from exc


@dataclass
class ResumenCargaSapDesglose:
    configurado: bool
    filas_leidas: int = 0
    filas_sin_clave: int = 0
    lineas_nuevas: int = 0
    lineas_actualizadas: int = 0
    lineas_sin_cambios: int = 0

    def to_dict(self) -> dict:
        return {
            "configurado": self.configurado,
            "filas_leidas": self.filas_leidas,
            "filas_sin_clave": self.filas_sin_clave,
            "lineas_nuevas": self.lineas_nuevas,
            "lineas_actualizadas": self.lineas_actualizadas,
            "lineas_sin_cambios": self.lineas_sin_cambios,
        }


def _texto(valor: object) -> Optional[str]:
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def _decimal(valor: object) -> Optional[Decimal]:
    if valor is None or valor == "":
        return None
    try:
        return Decimal(str(valor))
    except InvalidOperation:
        return None


def cargar_sap_desglose(db: Session, ruta_excel: Optional[str]) -> ResumenCargaSapDesglose:
    """Lee todas las hojas (mismo motivo que `cruce_codigos._cargar_indice`
    y `estado_sap.cargar_estado_sap`: no asumir que el dato vive siempre en
    la primera). Upsert por (`documento_compras`, `posicion`) -- la clave
    natural de una línea de pedido de compras real en SAP, nunca `id`, para
    que recargar el mismo fichero (o una exportación más reciente que
    incluya las mismas líneas) actualice en vez de duplicar."""
    resumen = ResumenCargaSapDesglose(configurado=bool(ruta_excel))
    if not ruta_excel:
        return resumen

    libro = openpyxl.load_workbook(ruta_excel, read_only=True, data_only=True)
    try:
        for hoja in libro.worksheets:
            filas_iter = hoja.iter_rows(values_only=True)
            primera = next(filas_iter, None)
            if primera is None:
                continue
            cabecera = [str(c).strip() if c is not None else "" for c in primera]
            if not all(col in cabecera for col in _COLUMNAS_REQUERIDAS):
                continue  # hoja sin las columnas mínimas: no es la hoja del desglose

            indice = {col: cabecera.index(col) for col in cabecera if col}

            def _valor(fila, columna):
                idx = indice.get(columna)
                return fila[idx] if idx is not None and idx < len(fila) else None

            for fila in filas_iter:
                resumen.filas_leidas += 1
                codigo = normalizar_codigo_expediente(_valor(fila, COLUMNA_EXPEDIENTE))
                documento_compras = _texto(_valor(fila, COLUMNA_DOCUMENTO_COMPRAS))
                posicion = _texto(_valor(fila, COLUMNA_POSICION))
                if not codigo or not documento_compras or not posicion:
                    resumen.filas_sin_clave += 1
                    continue

                datos = dict(
                    codigo_expediente=codigo,
                    material=_texto(_valor(fila, COLUMNA_MATERIAL)),
                    texto_breve=_texto(_valor(fila, COLUMNA_TEXTO_BREVE)),
                    cantidad_prevista=_decimal(_valor(fila, COLUMNA_CANTIDAD_PREVISTA)),
                    precio_neto=_decimal(_valor(fila, COLUMNA_PRECIO_NETO)),
                    unidad_medida=_texto(_valor(fila, COLUMNA_UNIDAD_MEDIDA)),
                )

                existente = db.execute(
                    select(SapDesgloseLinea).where(
                        SapDesgloseLinea.documento_compras == documento_compras,
                        SapDesgloseLinea.posicion == posicion,
                    )
                ).scalar_one_or_none()
                if existente is None:
                    db.add(SapDesgloseLinea(
                        documento_compras=documento_compras, posicion=posicion, **datos,
                    ))
                    resumen.lineas_nuevas += 1
                else:
                    cambio = any(getattr(existente, campo) != valor for campo, valor in datos.items())
                    for campo, valor in datos.items():
                        setattr(existente, campo, valor)
                    if cambio:
                        resumen.lineas_actualizadas += 1
                    else:
                        resumen.lineas_sin_cambios += 1
            db.commit()
    finally:
        libro.close()
    return resumen

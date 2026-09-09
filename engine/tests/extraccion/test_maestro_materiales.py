from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import select

from app.extraccion.maestro_materiales import (
    MaestroMaterialesPathInvalida,
    cargar_maestro_materiales,
    completar_unidades_desde_maestro,
    validar_ruta_maestro_materiales,
)
from app.models import (
    Documento,
    Expediente,
    EstadoExpediente,
    LineaCatalogo,
    Lote,
    MaestroMaterial,
    TipoDocumento,
)

_CABECERA = ["Material", "Texto breve", "Unidad medida base"]


@pytest.fixture
def excel_maestro(tmp_path):
    ruta = tmp_path / "maestro_materiales.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(_CABECERA)
    hoja.append(["603000210", "TRAVIESA DE ROBLE", "UN"])
    # Fila con espacios sobrantes (CONTEXTO.md sección 8): mismo criterio que
    # el resto de fuentes de entrada.
    hoja.append([" 611150412 ", "BRIDA PARA JUNTA ORDINARIA", "KG"])
    libro.save(ruta)
    return str(ruta)


def test_carga_crea_materiales_nuevos(db_session, excel_maestro):
    resumen = cargar_maestro_materiales(db_session, excel_maestro)

    assert resumen.configurado is True
    assert resumen.filas_leidas == 2
    assert resumen.materiales_nuevos == 2
    assert resumen.filas_sin_matricula == 0

    material = db_session.execute(
        select(MaestroMaterial).where(MaestroMaterial.matricula == "603000210")
    ).scalar_one()
    assert material.descripcion == "TRAVIESA DE ROBLE"
    assert material.unidad_medida == "UN"


def test_recorta_espacios_sobrantes_de_la_matricula(db_session, excel_maestro):
    cargar_maestro_materiales(db_session, excel_maestro)

    material = db_session.execute(
        select(MaestroMaterial).where(MaestroMaterial.matricula == "611150412")
    ).scalar_one_or_none()
    assert material is not None


def test_carga_repetida_actualiza_por_matricula_sin_duplicar(db_session, tmp_path):
    ruta = tmp_path / "v1.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(_CABECERA)
    hoja.append(["603000210", "TEXTO ANTIGUO", "UN"])
    libro.save(ruta)
    primera = cargar_maestro_materiales(db_session, str(ruta))

    ruta2 = tmp_path / "v2.xlsx"
    libro2 = openpyxl.Workbook()
    hoja2 = libro2.active
    hoja2.append(_CABECERA)
    hoja2.append(["603000210", "TEXTO CORREGIDO", "KG"])
    libro2.save(ruta2)
    segunda = cargar_maestro_materiales(db_session, str(ruta2))

    assert primera.materiales_nuevos == 1
    assert segunda.materiales_nuevos == 0
    assert segunda.materiales_actualizados == 1

    total = db_session.execute(select(MaestroMaterial)).scalars().all()
    assert len(total) == 1
    assert total[0].descripcion == "TEXTO CORREGIDO"
    assert total[0].unidad_medida == "KG"


def test_carga_repetida_identica_no_cuenta_como_actualizada(db_session, excel_maestro):
    cargar_maestro_materiales(db_session, excel_maestro)
    segunda = cargar_maestro_materiales(db_session, excel_maestro)

    assert segunda.materiales_nuevos == 0
    assert segunda.materiales_actualizados == 0
    assert segunda.materiales_sin_cambios == 2


def test_fila_sin_matricula_se_cuenta_y_se_ignora(db_session, tmp_path):
    ruta = tmp_path / "sin_matricula.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(_CABECERA)
    hoja.append([None, "SIN MATRICULA", "UN"])
    libro.save(ruta)

    resumen = cargar_maestro_materiales(db_session, str(ruta))

    assert resumen.filas_leidas == 1
    assert resumen.filas_sin_matricula == 1
    assert db_session.execute(select(MaestroMaterial)).scalar_one_or_none() is None


def test_sin_ruta_configurada_no_hace_nada(db_session):
    resumen = cargar_maestro_materiales(db_session, None)

    assert resumen.configurado is False
    assert resumen.filas_leidas == 0


# --- `validar_ruta_maestro_materiales` ---


def test_validar_ruta_sin_configurar_no_hace_nada():
    validar_ruta_maestro_materiales(None)
    validar_ruta_maestro_materiales("")


def test_validar_ruta_acepta_xlsx_valido(excel_maestro):
    validar_ruta_maestro_materiales(excel_maestro)


def test_validar_ruta_rechaza_ruta_inexistente(tmp_path):
    ruta = tmp_path / "no_existe.xlsx"

    with pytest.raises(MaestroMaterialesPathInvalida):
        validar_ruta_maestro_materiales(str(ruta))


# --- `completar_unidades_desde_maestro` ---


def _crear_linea(db_session, *, matricula, unidad_medida, precio=Decimal("10.00")):
    expediente = Expediente(
        codigo_expediente=f"6.24/28510.{matricula}",
        nombre_proyecto="Expediente de prueba",
        estado=EstadoExpediente.completado,
    )
    db_session.add(expediente)
    db_session.commit()
    lote = Lote(expediente_id=expediente.id, identificador_lote="1")
    db_session.add(lote)
    db_session.commit()
    documento = Documento(
        tipo_documento=TipoDocumento.anejo, hash=f"hash-{matricula}", ruta_almacenamiento=f"{matricula}.pdf"
    )
    db_session.add(documento)
    db_session.commit()
    linea = LineaCatalogo(
        lote_id=lote.id,
        expediente_id=expediente.id,
        clave_linea=f"clave-{matricula}",
        orden_aparicion=1,
        matricula=matricula,
        descripcion="Material de prueba",
        precio_unitario=precio,
        unidad_medida=unidad_medida,
        documento_origen_id=documento.id,
    )
    db_session.add(linea)
    db_session.commit()
    return linea


def test_completa_unidad_cuando_falta_y_hay_matricula(db_session, excel_maestro):
    cargar_maestro_materiales(db_session, excel_maestro)
    linea = _crear_linea(db_session, matricula="603000210", unidad_medida=None)

    resumen = completar_unidades_desde_maestro(db_session)

    assert resumen.lineas_evaluadas == 1
    assert resumen.lineas_completadas == 1
    db_session.refresh(linea)
    assert linea.unidad_medida == "UN"
    assert linea.unidad_medida_completada_desde_maestro is True


def test_no_pisa_unidad_ya_extraida_del_documento(db_session, excel_maestro):
    # El maestro dice "UN" pero el documento real ya trajo "KG" -- el
    # documento manda (CONTEXTO.md sección 12, y encargo explícito del
    # bloque 4: "sin sobrescribir lo extraído de los documentos").
    cargar_maestro_materiales(db_session, excel_maestro)
    linea = _crear_linea(db_session, matricula="603000210", unidad_medida="KG")

    resumen = completar_unidades_desde_maestro(db_session)

    assert resumen.lineas_evaluadas == 0
    assert resumen.lineas_completadas == 0
    db_session.refresh(linea)
    assert linea.unidad_medida == "KG"
    assert linea.unidad_medida_completada_desde_maestro is None


def test_no_completa_si_la_matricula_no_esta_en_el_maestro(db_session):
    linea = _crear_linea(db_session, matricula="999999999", unidad_medida=None)

    resumen = completar_unidades_desde_maestro(db_session)

    assert resumen.lineas_evaluadas == 1
    assert resumen.lineas_completadas == 0
    assert resumen.sin_matricula_en_maestro == 1
    db_session.refresh(linea)
    assert linea.unidad_medida is None


def test_sin_lineas_pendientes_no_hace_nada(db_session):
    resumen = completar_unidades_desde_maestro(db_session)

    assert resumen.lineas_evaluadas == 0
    assert resumen.lineas_completadas == 0

from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import select

from app.extraccion.sap_desglose import (
    SapDesglosePathInvalida,
    cargar_sap_desglose,
    validar_ruta_sap_desglose,
)
from app.models import SapDesgloseLinea

_CABECERA = [
    "expediente", "Documento compras", "Posición", "Indicador de borrado", "Última modificación",
    "Material", "Texto breve", "Centro", "Cantidad prevista", "Precio neto pedido", "Unidad medida base",
]


@pytest.fixture
def excel_sap_desglose(tmp_path):
    ruta = tmp_path / "contratos_traviesas.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(_CABECERA)
    hoja.append([
        "6.24/28510.0113", "3000000430", "10", "", None,
        "603000210", "TR-RB-2'60X0'24X0'14-54-PJ-1'668-1C-NEGR", "", 1, 117.06, "UN",
    ])
    # Fila con espacios sobrantes (CONTEXTO.md sección 8): mismo criterio que
    # cruce_codigos/estado_sap.
    hoja.append([
        " 6.24/28510.0114 ", "3000000431", "20", "", None,
        "603000204", "TR-RB-2'60X0'24X0'14-54-PI-1'688-1C-NEGR", "", 1, 117.06, "UN",
    ])
    libro.save(ruta)
    return str(ruta)


def test_carga_crea_lineas_nuevas(db_session, excel_sap_desglose):
    resumen = cargar_sap_desglose(db_session, excel_sap_desglose)

    assert resumen.configurado is True
    assert resumen.filas_leidas == 2
    assert resumen.lineas_nuevas == 2
    assert resumen.filas_sin_clave == 0

    linea = db_session.execute(
        select(SapDesgloseLinea).where(SapDesgloseLinea.documento_compras == "3000000430")
    ).scalar_one()
    assert linea.codigo_expediente == "6.24/28510.0113"
    assert linea.posicion == "10"
    assert linea.material == "603000210"
    assert linea.texto_breve == "TR-RB-2'60X0'24X0'14-54-PJ-1'668-1C-NEGR"
    assert linea.cantidad_prevista == Decimal("1")
    assert linea.precio_neto == Decimal("117.06")
    assert linea.unidad_medida == "UN"


def test_recorta_espacios_sobrantes_del_codigo_de_expediente(db_session, excel_sap_desglose):
    cargar_sap_desglose(db_session, excel_sap_desglose)

    linea = db_session.execute(
        select(SapDesgloseLinea).where(SapDesgloseLinea.documento_compras == "3000000431")
    ).scalar_one()
    assert linea.codigo_expediente == "6.24/28510.0114"


def test_carga_repetida_actualiza_por_documento_y_posicion_sin_duplicar(db_session, tmp_path):
    ruta = tmp_path / "v1.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(_CABECERA)
    hoja.append([
        "6.24/28510.0113", "3000000430", "10", "", None,
        "603000210", "TEXTO ANTIGUO", "", 1, 100.00, "UN",
    ])
    libro.save(ruta)
    primera = cargar_sap_desglose(db_session, str(ruta))

    ruta2 = tmp_path / "v2.xlsx"
    libro2 = openpyxl.Workbook()
    hoja2 = libro2.active
    hoja2.append(_CABECERA)
    hoja2.append([
        "6.24/28510.0113", "3000000430", "10", "", None,
        "603000210", "TEXTO CORREGIDO", "", 1, 117.06, "UN",
    ])
    libro2.save(ruta2)
    segunda = cargar_sap_desglose(db_session, str(ruta2))

    assert primera.lineas_nuevas == 1
    assert segunda.lineas_nuevas == 0
    assert segunda.lineas_actualizadas == 1

    total = db_session.execute(select(SapDesgloseLinea)).scalars().all()
    assert len(total) == 1
    assert total[0].texto_breve == "TEXTO CORREGIDO"
    assert total[0].precio_neto == Decimal("117.06")


def test_carga_repetida_identica_no_cuenta_como_actualizada(db_session, excel_sap_desglose):
    cargar_sap_desglose(db_session, excel_sap_desglose)
    segunda = cargar_sap_desglose(db_session, excel_sap_desglose)

    assert segunda.lineas_nuevas == 0
    assert segunda.lineas_actualizadas == 0
    assert segunda.lineas_sin_cambios == 2


def test_fila_sin_documento_de_compras_se_cuenta_y_se_ignora(db_session, tmp_path):
    ruta = tmp_path / "sin_documento.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(_CABECERA)
    hoja.append([
        "6.24/28510.0113", None, "10", "", None,
        "603000210", "SIN DOCUMENTO DE COMPRAS", "", 1, 117.06, "UN",
    ])
    libro.save(ruta)

    resumen = cargar_sap_desglose(db_session, str(ruta))

    assert resumen.filas_leidas == 1
    assert resumen.filas_sin_clave == 1
    assert db_session.execute(select(SapDesgloseLinea)).scalar_one_or_none() is None


def test_sin_ruta_configurada_no_hace_nada(db_session):
    resumen = cargar_sap_desglose(db_session, None)

    assert resumen.configurado is False
    assert resumen.filas_leidas == 0


# --- `validar_ruta_sap_desglose` ---


def test_validar_ruta_sin_configurar_no_hace_nada():
    validar_ruta_sap_desglose(None)
    validar_ruta_sap_desglose("")


def test_validar_ruta_acepta_xlsx_valido(excel_sap_desglose):
    validar_ruta_sap_desglose(excel_sap_desglose)


def test_validar_ruta_rechaza_ruta_inexistente(tmp_path):
    ruta = tmp_path / "no_existe.xlsx"

    with pytest.raises(SapDesglosePathInvalida):
        validar_ruta_sap_desglose(str(ruta))

import openpyxl
import pytest
from sqlalchemy import select

from app.extraccion.estado_sap import (
    EstadoSapPathInvalida,
    cargar_estado_sap,
    validar_ruta_estado_sap,
)
from app.models import Expediente


@pytest.fixture
def excel_estado_sap(tmp_path):
    ruta = tmp_path / "EXPEDIENTES_EJECUCION_SAP.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(["Expediente ADIF", "Título del expediente", "Descripción del estado"])
    hoja.append(["6.24/28510.0088", "SUMINISTRO DE BALASTO", "En ejecución"])
    # Fila con espacios sobrantes (CONTEXTO.md sección 8): un cruce exacto sin
    # normalizar fallaría en silencio, mismo hallazgo que cruce_codigos.
    hoja.append([" 6.14/28510.0126 ", "SUMINISTRO DE VISERAS", "En ejecución"])
    libro.save(ruta)
    return str(ruta)


def test_carga_crea_expedientes_nuevos_con_titulo_y_estado(db_session, excel_estado_sap):
    resumen = cargar_estado_sap(db_session, excel_estado_sap)

    assert resumen.configurado is True
    assert resumen.filas_leidas == 2
    assert resumen.expedientes_nuevos == 2
    assert resumen.expedientes_actualizados == 0

    nuevo = db_session.execute(
        select(Expediente).where(Expediente.codigo_expediente == "6.24/28510.0088")
    ).scalar_one()
    assert nuevo.nombre_proyecto == "SUMINISTRO DE BALASTO"
    assert nuevo.estado_contrato_sap == "En ejecución"
    assert nuevo.estado.value == "pendiente"  # estado de PROCESAMIENTO, sin tocar


def test_recorta_espacios_sobrantes_del_codigo(db_session, excel_estado_sap):
    cargar_estado_sap(db_session, excel_estado_sap)

    expediente = db_session.execute(
        select(Expediente).where(Expediente.codigo_expediente == "6.14/28510.0126")
    ).scalar_one()
    assert expediente.estado_contrato_sap == "En ejecución"


def test_no_pisa_el_nombre_de_proyecto_ya_extraido_de_un_pdf(db_session, excel_estado_sap):
    existente = Expediente(codigo_expediente="6.24/28510.0088", nombre_proyecto="Objeto real del PDF")
    db_session.add(existente)
    db_session.commit()

    cargar_estado_sap(db_session, excel_estado_sap)

    db_session.refresh(existente)
    assert existente.nombre_proyecto == "Objeto real del PDF"
    assert existente.estado_contrato_sap == "En ejecución"


def test_carga_repetida_no_duplica_ni_cuenta_como_nuevo(db_session, excel_estado_sap):
    primera = cargar_estado_sap(db_session, excel_estado_sap)
    segunda = cargar_estado_sap(db_session, excel_estado_sap)

    assert primera.expedientes_nuevos == 2
    assert segunda.expedientes_nuevos == 0
    assert segunda.expedientes_sin_cambios == 2

    total = db_session.execute(select(Expediente)).scalars().all()
    assert len(total) == 2


def test_carga_actualiza_estado_cuando_cambia_entre_exportaciones(db_session, tmp_path):
    ruta = tmp_path / "sap_v1.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(["Expediente ADIF", "Título del expediente", "Descripción del estado"])
    hoja.append(["6.24/28510.0088", "SUMINISTRO DE BALASTO", "En ejecución"])
    libro.save(ruta)
    cargar_estado_sap(db_session, str(ruta))

    ruta2 = tmp_path / "sap_v2.xlsx"
    libro2 = openpyxl.Workbook()
    hoja2 = libro2.active
    hoja2.append(["Expediente ADIF", "Título del expediente", "Descripción del estado"])
    hoja2.append(["6.24/28510.0088", "SUMINISTRO DE BALASTO", "Finalizado"])
    libro2.save(ruta2)
    resumen = cargar_estado_sap(db_session, str(ruta2))

    assert resumen.expedientes_actualizados == 1
    expediente = db_session.execute(
        select(Expediente).where(Expediente.codigo_expediente == "6.24/28510.0088")
    ).scalar_one()
    assert expediente.estado_contrato_sap == "Finalizado"


def test_sin_ruta_configurada_no_hace_nada(db_session):
    resumen = cargar_estado_sap(db_session, None)

    assert resumen.configurado is False
    assert resumen.filas_leidas == 0


def test_fila_sin_codigo_se_cuenta_y_se_ignora(db_session, tmp_path):
    ruta = tmp_path / "sap_fila_vacia.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(["Expediente ADIF", "Título del expediente", "Descripción del estado"])
    hoja.append([None, "Sin código", "En ejecución"])
    libro.save(ruta)

    resumen = cargar_estado_sap(db_session, str(ruta))

    assert resumen.filas_leidas == 1
    assert resumen.filas_sin_codigo == 1
    assert db_session.execute(select(Expediente)).scalar_one_or_none() is None


# --- `validar_ruta_estado_sap` (mismo criterio que
# `validar_ruta_codigos_proyecto`: fallar de forma visible al arrancar) ---


def test_validar_ruta_sin_configurar_no_hace_nada():
    validar_ruta_estado_sap(None)
    validar_ruta_estado_sap("")


def test_validar_ruta_acepta_xlsx_valido(excel_estado_sap):
    validar_ruta_estado_sap(excel_estado_sap)


def test_validar_ruta_rechaza_directorio(tmp_path):
    directorio = tmp_path / "estado_sap.xlsx"
    directorio.mkdir()

    with pytest.raises(EstadoSapPathInvalida):
        validar_ruta_estado_sap(str(directorio))


def test_validar_ruta_rechaza_ruta_inexistente(tmp_path):
    ruta = tmp_path / "no_existe.xlsx"

    with pytest.raises(EstadoSapPathInvalida):
        validar_ruta_estado_sap(str(ruta))

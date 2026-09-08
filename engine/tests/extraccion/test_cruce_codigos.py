from types import SimpleNamespace

import openpyxl
import pytest

from app.config import settings
from app.extraccion.cruce_codigos import (
    AutoreferenciaMatrizError,
    CodigosProyectoPathInvalida,
    asegurar_cruce_codigos,
    asignar_matriz,
    cruzar_codigo_proyecto,
    validar_ruta_codigos_proyecto,
)


@pytest.fixture
def excel_codigos(tmp_path):
    ruta = tmp_path / "Codigos_de_proyecto.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(["Nº Interno", "Nº Expediente", "MATRIZ", "ESPECIALIDAD/DISCIPLINA", "DESCRIPCIÓN"])
    hoja.append([24001, "6.24/28510.0128", None, "Señalización", "Equipos de medida"])
    # Fila con espacios sobrantes (CONTEXTO.md sección 8): un cruce exacto sin
    # normalizar fallaría en silencio.
    hoja.append([24038, "6.25/28510.0085 ", "6.25/28510.0028", "Vía", "Balasto"])
    libro.save(ruta)
    return str(ruta)


def test_cruza_por_numero_de_expediente_exacto(excel_codigos):
    resultado = cruzar_codigo_proyecto(excel_codigos, "6.24/28510.0128")

    assert resultado.cruzado is True
    assert resultado.codigo_interno == "24001"
    assert resultado.codigo_proyecto == "6.24/28510.0128"
    assert resultado.codigo_matriz is None


def test_recorta_espacios_sobrantes_antes_de_comparar(excel_codigos):
    resultado = cruzar_codigo_proyecto(excel_codigos, "6.25/28510.0085")

    assert resultado.cruzado is True
    assert resultado.codigo_interno == "24038"
    assert resultado.codigo_matriz == "6.25/28510.0028"


def test_cruza_por_matriz_cuando_no_hay_por_numero_de_expediente(excel_codigos):
    resultado = cruzar_codigo_proyecto(excel_codigos, "codigo-inexistente", codigo_matriz="6.25/28510.0028")

    assert resultado.cruzado is True
    assert resultado.codigo_interno == "24038"


def test_expediente_que_es_matriz_de_otra_fila_no_se_pone_su_propia_matriz(excel_codigos):
    # 6.25/28510.0028 nunca aparece como "Nº Expediente": solo es la MATRIZ
    # de la fila de 6.25/28510.0085 (acuerdo marco del que cuelga ese
    # pedido). Cruzar directamente por ese código (sección 3: "3 como
    # MATRIZ") debe traer su Nº Interno, pero NO debe devolver
    # codigo_matriz="6.25/28510.0028" (el propio código buscado, tal cual
    # sale de la columna MATRIZ de la fila encontrada) como si el expediente
    # tuviera una matriz distinta de sí mismo — eso sería inventarla por
    # auto-referencia (bug real: expediente 6.25/28510.0027, sesión de
    # pulido de la web, CONTEXTO.md).
    resultado = cruzar_codigo_proyecto(excel_codigos, "6.25/28510.0028")

    assert resultado.cruzado is True
    assert resultado.codigo_interno == "24038"
    assert resultado.codigo_matriz is None


def test_no_cruza_no_inventa_nada(excel_codigos):
    resultado = cruzar_codigo_proyecto(excel_codigos, "6.99/00000.0000")

    assert resultado.cruzado is False
    assert resultado.codigo_interno is None
    assert resultado.codigo_proyecto is None
    assert resultado.codigo_matriz is None


# --- `asignar_matriz`: único punto de escritura de `codigo_matriz`
# (docs/identidad-expediente.md, cuarta variante real de autorreferencia,
# caso 6.23/28510.0109) ---


def _expediente(**kwargs):
    base = dict(codigo_expediente=None, codigo_matriz=None, codigos_cruzados=None,
                codigo_interno=None, matriz_conflicto=None, codigo_matriz_en_cruce=None)
    base.update(kwargs)
    return SimpleNamespace(**base)


def test_asignar_matriz_escribe_cuando_no_hay_autorreferencia():
    expediente = _expediente(codigo_expediente="6.23/28510.0018")

    assert asignar_matriz(expediente, "6.23/28510.0102 ") is True
    assert expediente.codigo_matriz == "6.23/28510.0102"


def test_asignar_matriz_lanza_en_autorreferencia():
    expediente = _expediente(codigo_expediente="6.23/28510.0109")

    with pytest.raises(AutoreferenciaMatrizError):
        asignar_matriz(expediente, "6.23/28510.0109")
    assert expediente.codigo_matriz is None


def test_asignar_matriz_no_pisa_un_valor_existente_sin_pedirlo():
    expediente = _expediente(codigo_expediente="A", codigo_matriz="B")

    assert asignar_matriz(expediente, "C") is False
    assert expediente.codigo_matriz == "B"


def test_asignar_matriz_sobrescribe_si_se_pide_explicitamente():
    expediente = _expediente(codigo_expediente="A", codigo_matriz="B")

    assert asignar_matriz(expediente, "C", sobrescribir=True) is True
    assert expediente.codigo_matriz == "C"


def test_asignar_matriz_candidato_vacio_no_hace_nada():
    expediente = _expediente(codigo_expediente="A")

    assert asignar_matriz(expediente, None) is False
    assert asignar_matriz(expediente, "  ") is False
    assert expediente.codigo_matriz is None


# --- `asegurar_cruce_codigos`: caso real 6.23/28510.0109, licitación
# multi-lote cuya fila en el Excel de códigos declara su propia columna
# MATRIZ igual a su "Nº Expediente" -- ruido heredado del Excel (CONTEXTO.md,
# "Pendiente de resolver"), no un acuerdo marco real. Antes de esta sesión,
# `asegurar_cruce_codigos` escribía ese valor tal cual en
# `expediente.codigo_matriz`, dejando al expediente como su propia matriz. ---


def test_asegurar_cruce_codigos_matriz_autorreferenciada_en_excel_no_se_escribe(tmp_path, monkeypatch):
    ruta = tmp_path / "Codigos_de_proyecto.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(["Nº Interno", "Nº Expediente", "MATRIZ", "ESPECIALIDAD/DISCIPLINA", "DESCRIPCIÓN"])
    hoja.append([23026, "6.23/28510.0109", "6.23/28510.0109", "Vía", "SUMINISTRO DE BALASTO - 2 LOTES"])
    libro.save(ruta)
    monkeypatch.setattr(settings, "codigos_proyecto_path", str(ruta))

    expediente = _expediente(codigo_expediente="6.23/28510.0109")
    motivo = asegurar_cruce_codigos(None, expediente)

    assert motivo is None
    assert expediente.codigo_matriz is None
    assert expediente.codigo_interno == "23026"
    assert expediente.codigos_cruzados is True


# --- Reintento del cruce cuando cambia `codigo_matriz` (sesión de
# auditoría automática 2026-09-08, docs/sesion-2026-09-08-auditoria-
# automatica.md bloque 2 punto 2): un pedido derivado de acuerdo marco
# descubierto por el mecanismo inverso resuelve su `codigo_matriz` DESPUÉS
# de que el cruce ya se intentó una vez y falló -- caso real
# `6.26/28510.0032`/`0071`/`0014`. ---


def test_reintenta_el_cruce_cuando_aparece_una_matriz_nueva(excel_codigos, monkeypatch):
    monkeypatch.setattr(settings, "codigos_proyecto_path", excel_codigos)
    # Primer intento: sin matriz todavía, y un código propio que no cruza
    # por sí solo -- falla, como el primer intento real de un pedido recién
    # descubierto antes de que se le resuelva la matriz.
    expediente = _expediente(codigo_expediente="6.26/28510.9999", codigo_matriz=None)
    assert asegurar_cruce_codigos(None, expediente) is None
    assert expediente.codigos_cruzados is False
    assert expediente.codigo_interno is None
    assert expediente.codigo_matriz_en_cruce is None

    # El descubrimiento inverso matriz -> pedidos resuelve la matriz más
    # tarde, en un momento distinto -- el siguiente intento debe reintentar
    # el cruce (ahora sí encuentra la fila por MATRIZ) en vez de quedarse
    # con el `False` fijado la primera vez.
    expediente.codigo_matriz = "6.25/28510.0028"
    assert asegurar_cruce_codigos(None, expediente) is None
    assert expediente.codigos_cruzados is True
    assert expediente.codigo_interno == "24038"
    assert expediente.codigo_matriz_en_cruce == "6.25/28510.0028"


def test_no_reintenta_el_cruce_si_la_matriz_no_cambio(excel_codigos, monkeypatch):
    monkeypatch.setattr(settings, "codigos_proyecto_path", excel_codigos)
    expediente = _expediente(
        codigo_expediente="6.26/28510.9999", codigo_matriz="6.99/00000.0000",
        codigos_cruzados=False, codigo_matriz_en_cruce="6.99/00000.0000",
    )

    # Sin cambiar `codigo_matriz`, un segundo cruce real (que devolvería
    # otra cosa si se ejecutara) NUNCA debe intentarse: se apaga
    # `codigos_proyecto_path` para comprobar que la función no vuelve a
    # entrar en `cruzar_codigo_proyecto` -- si lo hiciera, fallaría por
    # ruta no configurada.
    monkeypatch.setattr(settings, "codigos_proyecto_path", None)
    assert asegurar_cruce_codigos(None, expediente) is None
    assert expediente.codigos_cruzados is False
    assert expediente.codigo_interno is None


def test_cruce_exitoso_nunca_se_reintenta_aunque_cambie_la_matriz(excel_codigos, monkeypatch):
    monkeypatch.setattr(settings, "codigos_proyecto_path", excel_codigos)
    expediente = _expediente(
        codigo_expediente="6.24/28510.0128", codigo_matriz=None,
        codigos_cruzados=True, codigo_interno="24001", codigo_matriz_en_cruce=None,
    )

    expediente.codigo_matriz = "cualquier-cosa-nueva"
    assert asegurar_cruce_codigos(None, expediente) is None
    assert expediente.codigo_interno == "24001"


# --- `validar_ruta_codigos_proyecto` (encargo de esta sesión: fallar de
# forma visible al arrancar en vez de dejar el cruce roto en silencio,
# como pasó con un bind-mount de Docker cuyo origen no existía). ---


def test_validar_ruta_codigos_proyecto_sin_configurar_no_hace_nada():
    validar_ruta_codigos_proyecto(None)
    validar_ruta_codigos_proyecto("")


def test_validar_ruta_codigos_proyecto_acepta_xlsx_valido(excel_codigos):
    validar_ruta_codigos_proyecto(excel_codigos)


def test_validar_ruta_codigos_proyecto_rechaza_directorio(tmp_path):
    # El caso real que motiva esto: un bind-mount de Docker cuyo origen no
    # existía, con Docker creando un directorio vacío en su lugar.
    directorio = tmp_path / "Codigos de proyecto.xlsx"
    directorio.mkdir()

    with pytest.raises(CodigosProyectoPathInvalida):
        validar_ruta_codigos_proyecto(str(directorio))


def test_validar_ruta_codigos_proyecto_rechaza_fichero_no_xlsx(tmp_path):
    ruta = tmp_path / "no_es_un_excel.xlsx"
    ruta.write_text("esto no es un .xlsx")

    with pytest.raises(CodigosProyectoPathInvalida):
        validar_ruta_codigos_proyecto(str(ruta))


def test_validar_ruta_codigos_proyecto_rechaza_ruta_inexistente(tmp_path):
    ruta = tmp_path / "no_existe.xlsx"

    with pytest.raises(CodigosProyectoPathInvalida):
        validar_ruta_codigos_proyecto(str(ruta))


# Aviso del cliente (sesión 2026-09-07, caso 6.26/28510.0014): antes solo se
# leía la primera hoja del Excel de códigos -- un expediente cuya única fila
# vive en una segunda hoja (p.ej. "en tramitación", separada de "en
# ejecución") nunca cruzaba.


def test_cruza_con_un_codigo_que_solo_vive_en_la_segunda_hoja(tmp_path):
    ruta = tmp_path / "Codigos_de_proyecto_multi_hoja.xlsx"
    libro = openpyxl.Workbook()
    hoja1 = libro.active
    hoja1.title = "En ejecución"
    hoja1.append(["Nº Interno", "Nº Expediente", "MATRIZ", "ESPECIALIDAD/DISCIPLINA", "DESCRIPCIÓN"])
    hoja1.append([24001, "6.24/28510.0128", None, "Señalización", "Equipos de medida"])
    hoja2 = libro.create_sheet("En tramitación")
    hoja2.append(["Nº Interno", "Nº Expediente", "MATRIZ", "ESPECIALIDAD/DISCIPLINA", "DESCRIPCIÓN"])
    hoja2.append([26014, "6.26/28510.0014", None, "Vía", "Pedido nº3 acuerdo marco de carril"])
    libro.save(ruta)

    resultado = cruzar_codigo_proyecto(str(ruta), "6.26/28510.0014")

    assert resultado.cruzado is True
    assert resultado.codigo_interno == "26014"


def test_hoja_vacia_no_revienta_la_lectura(tmp_path):
    ruta = tmp_path / "Codigos_de_proyecto_hoja_vacia.xlsx"
    libro = openpyxl.Workbook()
    hoja1 = libro.active
    hoja1.title = "En ejecución"
    hoja1.append(["Nº Interno", "Nº Expediente", "MATRIZ", "ESPECIALIDAD/DISCIPLINA", "DESCRIPCIÓN"])
    hoja1.append([24001, "6.24/28510.0128", None, "Señalización", "Equipos de medida"])
    libro.create_sheet("Hoja vacía")  # sin ninguna fila, ni cabecera
    libro.save(ruta)

    resultado = cruzar_codigo_proyecto(str(ruta), "6.24/28510.0128")

    assert resultado.cruzado is True
    assert resultado.codigo_interno == "24001"

"""Sesión 2026-09-21, bloque 1: el contraste de presupuestos (suma de cantidad
× precio de cada lote contra su presupuesto publicado), como hoja del Excel y
como vista de la web, y el lector de presupuestos por lote de los documentos.
"""
import io
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.contraste_presupuestos import (
    CIFRAS_DISTINTAS,
    CUADRA_AL_CENTIMO,
    CUADRA_MENOS_001,
    CUADRO_HEREDADO,
    EJECUCION_MATERIAL,
    FALTAN_CANTIDADES,
    INCIERTO,
    NO_CUADRA_DOCUMENTO,
    NO_CUADRA_FALTAN_LINEAS,
    NO_CUADRA_SIN_CAUSA,
    NOMBRE_HOJA,
    RESULTADOS,
    SIN_PRESUPUESTO,
    SOLO_TOTAL_DEL_EXPEDIENTE,
    construir_contraste,
    lineas_de_materiales_por_lote,
)
from app.db import Base, get_db
from app.extraccion.presupuesto_lote import (
    ANUNCIO,
    CONTRATO,
    EJECUCION_MATERIAL as REDACCION_EJECUCION_MATERIAL,
    LISTA_SIN_IVA,
    leer_presupuestos_de_lote,
    registrar_presupuestos_de_lote,
)
from app.extraccion.texto import PaginaTexto
from app.models import (
    Documento,
    DocumentoExpediente,
    EstadoExpediente,
    Expediente,
    LineaCatalogo,
    Lote,
    PresupuestoLoteDocumento,
    SindicacionExpediente,
    TipoDocumento,
    TrazaOrigen,
)


# --- El lector de presupuestos por lote -------------------------------------


def test_el_bloque_de_lote_del_anuncio_da_las_dos_cifras():
    texto = (
        "Suministro de traviesas. 6 Lotes. Lote 3: Área Territorial Norte. Presupuesto base de "
        "licitación Lugar de ejecución Importe 9.922.000 EUR. País ES Importe (sin impuestos) "
        "8.200.000 EUR. Subentidad Nacional España"
    )
    [leido] = leer_presupuestos_de_lote([PaginaTexto(numero=4, texto=texto)])
    assert (leido.identificador_lote, leido.redaccion, leido.pagina) == ("3", ANUNCIO, 4)
    assert leido.importe == Decimal("8200000")
    assert leido.importe_con_iva == Decimal("9922000")


def test_una_mencion_de_lote_no_se_liga_a_una_cifra_si_hay_otra_mencion_en_medio():
    # La cabecera de un anuncio que enumera los lotes antes del presupuesto
    # GLOBAL: ese total no es del lote 1 ni del 2.
    texto = (
        "Objeto: Lote 1: Norte; Lote 2: Sur. Presupuesto base de licitación Importe 1.210.000 EUR. "
        "Importe (sin impuestos) 1.000.000 EUR."
    )
    leidos = leer_presupuestos_de_lote([PaginaTexto(numero=1, texto=texto)])
    assert [(l.identificador_lote, l.importe) for l in leidos] == [("2", Decimal("1000000"))]


def test_el_numero_de_lote_con_ceros_es_el_mismo_lote():
    texto = (
        "Nº Lote: 003 Objeto del Contrato: Cerrojos. Presupuesto base de licitación Importe "
        "1.403.600 EUR. Importe (sin impuestos) 1.160.000 EUR."
    )
    [leido] = leer_presupuestos_de_lote([PaginaTexto(numero=2, texto=texto)])
    assert leido.identificador_lote == "3"


def test_contrato_lista_sin_iva_y_ejecucion_material():
    texto = (
        "Ascendiendo el importe de licitación del lote 2 a 7.400.000,00 € (IVA excluido). "
        "Lote 5: Área Territorial Este, 5.400.000,00 € sin IVA; "
        "Presupuesto de Ejecución Material 2.975.673,96 €"
    )
    leidos = {l.redaccion: l for l in leer_presupuestos_de_lote([PaginaTexto(numero=9, texto=texto)])}
    assert (leidos[CONTRATO].identificador_lote, leidos[CONTRATO].importe) == ("2", Decimal("7400000.00"))
    assert (leidos[LISTA_SIN_IVA].identificador_lote, leidos[LISTA_SIN_IVA].importe) == (
        "5", Decimal("5400000.00")
    )
    assert leidos[REDACCION_EJECUCION_MATERIAL].identificador_lote is None
    assert leidos[REDACCION_EJECUCION_MATERIAL].importe == Decimal("2975673.96")


def _expediente(db, codigo="6.20/28510.0054", importe=None):
    expediente = Expediente(
        codigo_expediente=codigo, estado=EstadoExpediente.completado, importe_licitacion=importe
    )
    db.add(expediente)
    db.flush()
    return expediente


def _documento(db, expediente, nombre="PLIEGO_1.pdf", hash_="h1"):
    documento = Documento(hash=hash_, tipo_documento=TipoDocumento.pliego, ruta_almacenamiento=f"/x/{hash_}")
    db.add(documento)
    db.flush()
    db.add(DocumentoExpediente(documento_id=documento.id, expediente_id=expediente.id, nombre_archivo=nombre))
    db.flush()
    return documento


def test_registrar_reescribe_las_apariciones_del_expediente(db_session):
    expediente = _expediente(db_session)
    documento = _documento(db_session, expediente)
    pagina = PaginaTexto(numero=1, texto="Lote 2: Sur, 100.000,00 € sin IVA")
    assert registrar_presupuestos_de_lote(db_session, expediente.id, [(documento.id, [pagina])]) == 1
    db_session.commit()
    # Una segunda pasada no duplica: borra y vuelve a escribir.
    registrar_presupuestos_de_lote(db_session, expediente.id, [(documento.id, [pagina])])
    db_session.commit()
    assert len(db_session.execute(select(PresupuestoLoteDocumento)).scalars().all()) == 1
    # Y una pasada sin documentos vacía lo de la anterior.
    registrar_presupuestos_de_lote(db_session, expediente.id, [])
    db_session.commit()
    assert db_session.execute(select(PresupuestoLoteDocumento)).scalars().all() == []


# --- El contraste --------------------------------------------------------------


def _lote(db, expediente, identificador="1", importe=None, codigo_lote=None):
    lote = Lote(
        expediente_id=expediente.id, identificador_lote=identificador, importe_licitacion=importe,
        codigo_expediente_lote=codigo_lote,
    )
    db.add(lote)
    db.flush()
    return lote


def _linea(db, expediente, lote, cantidad, precio, descripcion="material", heredada=None, orden=[0]):
    orden[0] += 1
    linea = LineaCatalogo(
        expediente_id=expediente.id, lote_id=lote.id, clave_linea=f"k{orden[0]}", orden_aparicion=orden[0],
        descripcion=descripcion, cantidad=None if cantidad is None else Decimal(cantidad),
        precio_unitario=None if precio is None else Decimal(precio), heredado_de_matriz=heredada,
    )
    db.add(linea)
    db.flush()
    return linea


def _traza(db, entidad_tipo, entidad_id, documento, fragmento, pagina=1):
    db.add(TrazaOrigen(
        entidad_tipo=entidad_tipo, entidad_id=entidad_id, campo="importe_licitacion",
        documento_id=documento.id, pagina=pagina, fragmento=fragmento,
    ))
    db.flush()


def _contraste(db):
    db.commit()
    return construir_contraste(db, lineas_de_materiales_por_lote(db))


def test_cuadra_al_centimo_con_la_base_sin_iva_del_contrato(db_session):
    expediente = _expediente(db_session)
    documento = _documento(db_session, expediente, "CONTRATO_1.pdf")
    lote = _lote(db_session, expediente, "3", importe=Decimal("1000.00"), codigo_lote="6.20/28510.0056")
    _lote(db_session, expediente, "4")
    _traza(db_session, "lote", lote.id, documento,
           "importe de licitación del Lote 3 a 1.000,00 € (IVA excluido)", pagina=103)
    _linea(db_session, expediente, lote, "10", "40")
    _linea(db_session, expediente, lote, "2", "300")
    [fila] = _contraste(db_session).filas
    assert fila.resultado == CUADRA_AL_CENTIMO
    assert fila.tipo_cifra == "Presupuesto base de licitación sin IVA"
    assert (fila.documento, fila.pagina) == ("CONTRATO_1.pdf", 103)
    assert fila.suma_lineas == fila.cifra_comparada == Decimal("1000.00")
    assert (fila.lineas, fila.lineas_sin_cantidad, fila.diferencia) == (2, 0, Decimal("0.00"))


def test_ejecucion_material_declarada_con_el_centimo_del_redondeo(db_session):
    """El caso de `6.17/28510.0056` lote 1: el cuadro suma su ejecución
    material, y 2.975.673,96 + 9 % + 6 % redondeados por separado es el
    3.422.025,06 que publica la Plataforma (×1,15 de una vez daría ,05)."""
    expediente = _expediente(db_session, "6.17/28510.0056")
    documento = _documento(db_session, expediente, "ANEJO_1.pdf")
    lote = _lote(db_session, expediente, "1")
    _lote(db_session, expediente, "2")
    db_session.add(SindicacionExpediente(
        codigo_expediente="6.17/28510.0056", expediente_id=expediente.id, periodo_zip="202412",
        actualizado_en=__import__("datetime").datetime(2024, 12, 1),
        lotes=[{"identificador": "001", "importe_licitacion_sin_impuestos": "3422025.06",
                "importe_licitacion_con_impuestos": "4140650.32"}],
    ))
    db_session.add(PresupuestoLoteDocumento(
        expediente_id=expediente.id, identificador_lote=None, importe=Decimal("2975673.96"),
        redaccion=REDACCION_EJECUCION_MATERIAL, documento_id=documento.id, pagina=42,
        fragmento="Presupuesto de Ejecución Material 2.975.673,96 €",
    ))
    _linea(db_session, expediente, lote, "1", "2975673.96")
    [fila] = _contraste(db_session).filas
    assert fila.resultado == CUADRA_AL_CENTIMO
    assert fila.tipo_cifra == EJECUCION_MATERIAL
    assert fila.presupuesto_publicado == Decimal("3422025.06")
    assert fila.cifra_comparada == Decimal("2975673.96")
    assert fila.documento.startswith("Sindicación de la Plataforma (boletín 12/2024)")
    assert "ANEJO_1.pdf p.42" in fila.explicacion
    assert "Comprobado a mano" in fila.explicacion


def test_sin_ejecucion_material_declarada_se_compara_con_la_base_entre_115(db_session):
    expediente = _expediente(db_session, "3.25/28510.0164", importe=Decimal("74750.00"))
    documento = _documento(db_session, expediente, "ADJUDICACION_1.pdf")
    lote = _lote(db_session, expediente, "1", importe=Decimal("74750.00"))
    _traza(db_session, "expediente", expediente.id, documento, "Presupuesto de licitación: 74.750,00 €")
    _linea(db_session, expediente, lote, "1", "65000")
    [fila] = _contraste(db_session).filas
    assert (fila.resultado, fila.tipo_cifra) == (CUADRA_AL_CENTIMO, EJECUCION_MATERIAL)
    assert fila.cifra_comparada == Decimal("65000.00")
    assert "Ningún documento declara esa ejecución material" in fila.explicacion


def test_faltan_cantidades_y_la_partida_alzada_cuenta_una_vez(db_session):
    expediente = _expediente(db_session, importe=Decimal("1000"))
    documento = _documento(db_session, expediente, "ADJUDICACION_1.pdf")
    lote = _lote(db_session, expediente, "1", importe=Decimal("1000"))
    _traza(db_session, "expediente", expediente.id, documento,
           "Presupuesto base de licitación Importe 1.210 EUR. Importe (sin impuestos) 1.000 EUR")
    _linea(db_session, expediente, lote, "9", "100")
    _linea(db_session, expediente, lote, None, "100", descripcion="Partida alzada a justificar")
    [fila] = _contraste(db_session).filas
    assert fila.resultado == CUADRA_AL_CENTIMO
    assert "partida(s) alzada(s)" in fila.explicacion
    _linea(db_session, expediente, lote, None, "5")
    [fila] = _contraste(db_session).filas
    assert (fila.resultado, fila.lineas_sin_cantidad) == (FALTAN_CANTIDADES, 1)


def test_menos_del_001_y_no_cuadra_con_el_porcentaje_exacto(db_session):
    a = _expediente(db_session, "6.24/28510.0001", importe=Decimal("1000000"))
    doc_a = _documento(db_session, a, "ADJUDICACION_1.pdf", "ha")
    lote_a = _lote(db_session, a, "1", importe=Decimal("1000000"))
    _traza(db_session, "expediente", a.id, doc_a, "Presupuesto de licitación: 1.000.000,00 €")
    _linea(db_session, a, lote_a, "1", "999999.97")
    b = _expediente(db_session, "6.24/28510.0002", importe=Decimal("2400000"))
    doc_b = _documento(db_session, b, "ADJUDICACION_1.pdf", "hb")
    lote_b = _lote(db_session, b, "1", importe=Decimal("2400000"))
    _traza(db_session, "expediente", b.id, doc_b, "Presupuesto de licitación: 2.400.000,00 €")
    _linea(db_session, b, lote_b, "1", "2160000")
    filas = {f.codigo_expediente: f for f in _contraste(db_session).filas}
    assert filas["6.24/28510.0001"].resultado == CUADRA_MENOS_001
    assert filas["6.24/28510.0002"].resultado == NO_CUADRA_SIN_CAUSA
    assert "exactamente el 90 %" in filas["6.24/28510.0002"].explicacion


def test_la_discrepancia_comprobada_a_mano_lleva_lo_que_se_comprobo(db_session):
    expediente = _expediente(db_session, "6.20/28510.0054")
    documento = _documento(db_session, expediente)
    lote = _lote(db_session, expediente, "3")
    _lote(db_session, expediente, "4")
    db_session.add(PresupuestoLoteDocumento(
        expediente_id=expediente.id, identificador_lote="3", importe=Decimal("8200000"),
        importe_con_iva=Decimal("9922000"), redaccion=ANUNCIO, documento_id=documento.id, pagina=4,
    ))
    linea = _linea(db_session, expediente, lote, "1", "8230002.13")
    [fila] = _contraste(db_session).filas
    assert fila.resultado == NO_CUADRA_DOCUMENTO
    assert "ANEJO_8.pdf" in fila.explicacion
    # Con otras cifras, la causa ya no se escribe: explicaría otra cosa.
    linea.precio_unitario = Decimal("8230002.14")
    [fila] = _contraste(db_session).filas
    assert fila.resultado == NO_CUADRA_SIN_CAUSA
    assert "ANEJO_8.pdf" not in fila.explicacion


def test_el_lote_que_pone_el_sistema_no_toma_el_presupuesto_del_lote_1_de_un_documento(db_session):
    expediente = _expediente(db_session)
    documento = _documento(db_session, expediente)
    lote = _lote(db_session, expediente, "1")  # único y sin código propio: es el del sistema
    db_session.add(PresupuestoLoteDocumento(
        expediente_id=expediente.id, identificador_lote="1", importe=Decimal("500"),
        redaccion=LISTA_SIN_IVA, documento_id=documento.id, pagina=1,
    ))
    _linea(db_session, expediente, lote, "5", "100")
    contraste = _contraste(db_session)
    assert contraste.filas == []
    assert [f.motivo for f in contraste.fuera] == [SIN_PRESUPUESTO]


def test_los_lotes_que_no_entran_llevan_su_motivo(db_session):
    # Cuadro heredado del acuerdo marco.
    pedido = _expediente(db_session, "6.24/28510.0040", importe=Decimal("100"))
    _linea(db_session, pedido, _lote(db_session, pedido, "1", importe=Decimal("100")), "1", "7", heredada=True)
    # Cifras distintas para el mismo lote.
    dos = _expediente(db_session, "6.20/28510.0115")
    documento = _documento(db_session, dos, "CONTRATO_1.pdf", "h2")
    lote_dos = _lote(db_session, dos, "1")
    _lote(db_session, dos, "2")
    for importe in ("15000", "35000"):
        db_session.add(PresupuestoLoteDocumento(
            expediente_id=dos.id, identificador_lote="1", importe=Decimal(importe),
            redaccion=CONTRATO, documento_id=documento.id, pagina=1,
        ))
    _linea(db_session, dos, lote_dos, "1", "10")
    # Varios lotes y solo el total del expediente.
    total = _expediente(db_session, "6.25/28510.0019", importe=Decimal("4950000"))
    lote_total = _lote(db_session, total, "1")
    _lote(db_session, total, "2")
    _linea(db_session, total, lote_total, "1", "10")
    motivos = {f.codigo_expediente: f.motivo for f in _contraste(db_session).fuera}
    assert motivos == {
        "6.24/28510.0040": CUADRO_HEREDADO,
        "6.20/28510.0115": CIFRAS_DISTINTAS,
        "6.25/28510.0019": SOLO_TOTAL_DEL_EXPEDIENTE,
    }


def test_una_etiqueta_que_no_dice_si_lleva_iva_se_avisa_en_la_fila(db_session):
    expediente = _expediente(db_session, importe=Decimal("100"))
    documento = _documento(db_session, expediente, "ADJUDICACION_1.pdf")
    lote = _lote(db_session, expediente, "1", importe=Decimal("100"))
    _traza(db_session, "expediente", expediente.id, documento, "Presupuesto máximo: 100,00 €")
    _linea(db_session, expediente, lote, "1", "100")
    [fila] = _contraste(db_session).filas
    assert fila.tipo_cifra.startswith(INCIERTO)
    assert fila.resultado == CUADRA_AL_CENTIMO


# --- La hoja del Excel y la vista de la web dan lo mismo ---------------------


@pytest.fixture()
def db_api():
    from app import models as m

    engine = create_engine(
        "sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine, tables=[
        m.Expediente.__table__, m.Lote.__table__, m.Documento.__table__,
        m.DocumentoExpediente.__table__, m.LineaCatalogo.__table__, m.TrazaOrigen.__table__,
        m.SindicacionExpediente.__table__, m.CacheOcrDocumento.__table__, m.TrabajoCola.__table__,
        m.PresupuestoLoteDocumento.__table__,
    ])
    sesion = sessionmaker(bind=engine)()
    try:
        yield sesion
    finally:
        sesion.close()


@pytest.fixture()
def cliente(db_api):
    from app.main import app

    def _get_db_prueba():
        yield db_api

    app.dependency_overrides[get_db] = _get_db_prueba
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.clear()


def _dos_lotes(db):
    expediente = _expediente(db, importe=Decimal("1000"))
    documento = _documento(db, expediente, "CONTRATO_1.pdf")
    uno = _lote(db, expediente, "1", importe=Decimal("1000"))
    dos = _lote(db, expediente, "2", importe=Decimal("500"))
    _traza(db, "lote", uno.id, documento, "importe de licitación del lote 1 a 1.000,00 € (IVA excluido)")
    _traza(db, "lote", dos.id, documento, "importe de licitación del lote 2 a 500,00 € (IVA excluido)")
    _linea(db, expediente, uno, "10", "100")
    _linea(db, expediente, dos, "1", "400")
    db.commit()


def test_la_api_devuelve_las_filas_su_recuento_y_filtra(cliente, db_api):
    _dos_lotes(db_api)
    cuerpo = cliente.get("/contraste-presupuestos").json()
    assert cuerpo["total"] == 2
    assert [r["resultado"] for r in cuerpo["resultados"]] == list(RESULTADOS)
    assert all(r["significado"] for r in cuerpo["resultados"])
    assert {r["resultado"]: r["lotes"] for r in cuerpo["resultados"]}[CUADRA_AL_CENTIMO] == 1
    filtrado = cliente.get("/contraste-presupuestos", params={"resultado": NO_CUADRA_SIN_CAUSA}).json()
    assert [f["lote"] for f in filtrado["filas"]] == ["2"]
    assert filtrado["total"] == 2  # el recuento es el del total, no el del filtro


def test_la_hoja_del_excel_da_las_mismas_filas_que_la_api(cliente, db_api):
    from app.exportacion import generar_excel_catalogo

    _dos_lotes(db_api)
    api = cliente.get("/contraste-presupuestos").json()["filas"]
    libro = load_workbook(io.BytesIO(generar_excel_catalogo(db_api)))
    assert libro.sheetnames.index(NOMBRE_HOJA) == libro.sheetnames.index("Conciliación") + 1
    hoja = libro[NOMBRE_HOJA]
    filas = [f for f in hoja.iter_rows(min_row=2, values_only=True) if f[0] and f[1] and f[12]][: len(api)]
    assert [(f[0], f[1], f[12]) for f in filas] == [
        (a["codigo_expediente"], a["lote"], a["resultado"]) for a in api
    ]
    assert [Decimal(str(f[7])) for f in filas] == [Decimal(a["suma_lineas"]) for a in api]


def test_una_cifra_guardada_sin_traza_se_demuestra_con_el_documento_que_dice_la_misma(db_session):
    expediente = _expediente(db_session, "6.25/28510.0027")
    documento = _documento(db_session, expediente, "PLIEGO_1.pdf")
    lote = _lote(db_session, expediente, "1", importe=Decimal("593375.00"))
    _lote(db_session, expediente, "3", importe=Decimal("853250.00"))
    db_session.add(PresupuestoLoteDocumento(
        expediente_id=expediente.id, identificador_lote="1", importe=Decimal("593375"),
        importe_con_iva=Decimal("717983.75"), redaccion=ANUNCIO, documento_id=documento.id, pagina=5,
    ))
    _linea(db_session, expediente, lote, "1", "593375")
    [fila] = _contraste(db_session).filas
    assert fila.tipo_cifra == "Presupuesto base de licitación sin IVA"
    assert (fila.documento, fila.pagina) == ("PLIEGO_1.pdf", 5)
    # Si el documento dijera otra cifra, se queda la guardada y se avisa.
    db_session.execute(
        __import__("sqlalchemy").update(PresupuestoLoteDocumento).values(importe=Decimal("600000"))
    )
    [fila] = _contraste(db_session).filas
    assert fila.tipo_cifra.startswith(INCIERTO)
    assert fila.presupuesto_publicado == Decimal("593375.00")


def test_las_filas_del_lote_que_no_salen_en_materiales_se_dicen(db_session):
    """`6.25/28510.0097` lote 2: sus filas sin cantidad no salen en
    "Materiales" (mapeo incoherente), así que no se suman; la fila lo dice."""
    from app.catalogo import MOTIVO_MAPEO_INCOHERENTE

    expediente = _expediente(db_session, importe=Decimal("1000"))
    documento = _documento(db_session, expediente, "ADJUDICACION_1.pdf")
    lote = _lote(db_session, expediente, "1", importe=Decimal("1000"))
    _traza(db_session, "expediente", expediente.id, documento, "Presupuesto de licitación: 1.000,00 €")
    _linea(db_session, expediente, lote, "5", "100")
    fuera = _linea(db_session, expediente, lote, None, "10")
    fuera.motivo_revision = MOTIVO_MAPEO_INCOHERENTE
    [fila] = _contraste(db_session).filas
    assert (fila.lineas, fila.lineas_sin_cantidad, fila.resultado) == (1, 0, NO_CUADRA_FALTAN_LINEAS)
    assert "1 fila(s) de este lote están en revisión" in fila.explicacion

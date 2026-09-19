"""Bloque 2, sesión 2026-09-19 (sexta parte): las tres entradas que el cliente
tiene que mandarnos, con su ingesta, su cruce y su salida montados y probados
**sobre datos sintéticos** -- los ficheros reales no han llegado.

Cada prueba cubre las dos mitades del encargo: qué hace el sistema con el
fichero en la forma esperada, y qué hace si llega con otra.
"""
from decimal import Decimal

import openpyxl
import pytest
from openpyxl import Workbook

from app.catalogo_antiguo import (
    EMPAREJADO_POR_DESCRIPCION,
    EMPAREJADO_POR_MATRICULA,
    InformeCatalogoAntiguo,
    MaterialCatalogoAntiguo,
    MaterialNuestro,
    ResumenCatalogoAntiguo,
    clave_descripcion,
    generar_informe_catalogo_antiguo,
    leer_catalogo_antiguo,
)
from app.conciliacion import APORTA_LINEAS, FilaConciliacion
from app.extraccion.estados_adif import _importe, _indice_presupuesto
from app.extraccion.vigentes_remanente import (
    ACCION_ALTA_Y_BUSCADO,
    ACCION_EN_LA_CONCILIACION,
    ACCION_SIN_BUSCAR,
    cruzar_vigentes_con_remanente,
    leer_codigos_vigentes,
)
from app.models import EstadoExpediente, Expediente, TrabajoCola


def _libro(filas, nombre="Hoja1", ruta=None):
    libro = Workbook()
    hoja = libro.active
    hoja.title = nombre
    for fila in filas:
        hoja.append(list(fila))
    libro.save(ruta)
    return str(ruta)


# --------------------------------------------------------------------------
# 1 — EXPEDIENTES_VIGENTES_CON_REMANENTE.xlsx
# --------------------------------------------------------------------------


def test_la_lista_de_vigentes_se_lee_por_el_encabezado(tmp_path):
    ruta = _libro(
        [
            ("Expediente ADIF", "Remanente"),
            ("6.24/28510.0088", 12345.0),
            ("3.21/28510.0096", 900.0),
            (None, None),
        ],
        ruta=tmp_path / "vigentes.xlsx",
    )
    codigos, localizada_por, hoja, filas = leer_codigos_vigentes(ruta)
    assert codigos == ["6.24/28510.0088", "3.21/28510.0096"]
    assert localizada_por == "encabezado"
    assert (hoja, filas) == ("Hoja1", 3)


def test_sin_encabezado_reconocible_se_localiza_por_la_forma_de_los_codigos(tmp_path):
    """"De formato desconocido" no es excusa para adivinar: se acepta una
    columna sin encabezado solo si la mayoría de sus valores tienen forma de
    código de expediente de ADIF, y solo si es la única así."""
    ruta = _libro(
        [
            ("Vigentes con remanente 2026", None),
            ("6.24/28510.0088", 12345.0),
            ("3.21/28510.0096", 900.0),
            ("28510/2023", 10.0),
        ],
        ruta=tmp_path / "sin_encabezado.xlsx",
    )
    codigos, localizada_por, _hoja, _filas = leer_codigos_vigentes(ruta)
    assert localizada_por == "contenido"
    assert codigos == ["6.24/28510.0088", "3.21/28510.0096", "28510/2023"]


def test_un_fichero_sin_ninguna_columna_de_codigos_no_se_inventa(tmp_path):
    ruta = _libro(
        [("Concepto", "Importe"), ("Material A", 10.0), ("Material B", 20.0)],
        ruta=tmp_path / "otro.xlsx",
    )
    codigos, localizada_por, hoja, _filas = leer_codigos_vigentes(ruta)
    assert (codigos, localizada_por, hoja) == ([], "no encontrada", None)


def test_el_cruce_da_la_situacion_de_la_conciliacion_y_busca_lo_que_falta(db_session, tmp_path):
    existente = Expediente(codigo_expediente="6.24/28510.0088", estado=EstadoExpediente.completado)
    db_session.add(existente)
    db_session.commit()

    ruta = _libro(
        [
            ("Expediente", "Remanente"),
            ("6.24/28510.0088", 1.0),
            ("9.99/28510.9999", 2.0),
        ],
        ruta=tmp_path / "vigentes.xlsx",
    )
    conciliacion = [
        FilaConciliacion(
            codigo_expediente="6.24/28510.0088", titulo="t", organo_contratacion=None,
            estado_plataforma=None, estado_adif=None, documentos_descargados=3,
            documentos_reconocimiento_optico=0, lineas_en_catalogo=12, baja="0,50 %",
            situacion=APORTA_LINEAS, motivo="",
        )
    ]
    resumen = cruzar_vigentes_con_remanente(db_session, ruta, conciliacion, buscar=True)

    assert resumen.formato_reconocido is True
    assert (resumen.en_la_conciliacion, resumen.fuera_de_la_conciliacion) == (1, 1)
    assert resumen.por_situacion == {APORTA_LINEAS: 1}
    assert resumen.dados_de_alta == 1 and resumen.busquedas_encoladas == 1
    acciones = {f.codigo_expediente: f.accion for f in resumen.filas}
    assert acciones["6.24/28510.0088"] == ACCION_EN_LA_CONCILIACION
    assert acciones["9.99/28510.9999"] == ACCION_ALTA_Y_BUSCADO
    # La búsqueda es la vía normal del sistema: un trabajo de la cola.
    trabajos = db_session.query(TrabajoCola).filter(TrabajoCola.tipo == "descargar_expediente").all()
    assert len(trabajos) == 1


def test_el_cruce_sin_buscar_no_da_de_alta_ni_encola_nada(db_session, tmp_path):
    ruta = _libro(
        [("Expediente",), ("9.99/28510.9999",)], ruta=tmp_path / "vigentes.xlsx"
    )
    resumen = cruzar_vigentes_con_remanente(db_session, ruta, [], buscar=False)
    assert resumen.fuera_de_la_conciliacion == 1
    assert (resumen.dados_de_alta, resumen.busquedas_encoladas) == (0, 0)
    assert resumen.filas[0].accion == ACCION_SIN_BUSCAR
    assert db_session.query(Expediente).count() == 0
    assert db_session.query(TrabajoCola).count() == 0


def test_sin_ruta_configurada_el_cruce_no_hace_nada(db_session):
    resumen = cruzar_vigentes_con_remanente(db_session, None, [], buscar=True)
    assert resumen.configurado is False and resumen.filas == []


# --------------------------------------------------------------------------
# 2 — El listado de estados de ADIF con presupuesto de licitación
# --------------------------------------------------------------------------


def test_la_columna_de_presupuesto_se_reconoce_por_el_nombre():
    for nombre in ("Presupuesto", "Presupuesto de licitación", "PRESUPUESTO BASE DE LICITACIÓN",
                   "Importe de licitación", "PBL"):
        assert _indice_presupuesto(["Expediente ADIF", nombre]) == 1, nombre


def test_un_importe_de_adjudicacion_no_es_el_presupuesto_de_licitacion():
    """Contrastar el presupuesto contra un importe adjudicado compararía dos
    cosas distintas y daría diferencias en casi todas las filas."""
    assert _indice_presupuesto(["Expediente ADIF", "Importe de adjudicación"]) is None
    assert _indice_presupuesto(["Expediente ADIF", "Presupuesto con IVA"]) is None


def test_con_dos_columnas_candidatas_no_se_adivina():
    assert _indice_presupuesto(["Presupuesto", "Presupuesto base"]) is None


def test_el_listado_de_hoy_no_trae_ninguna():
    """Las cuatro columnas reales del fichero del 18/09/2026."""
    assert _indice_presupuesto(
        ["Título del expediente", "Expediente ADIF", "Fecha de creación", "Descripción del estado"]
    ) is None


def test_el_presupuesto_se_lee_como_numero_y_como_texto_en_formato_espanol():
    assert _importe(1234567.89) == Decimal("1234567.89")
    assert _importe("1.234.567,89 €") == Decimal("1234567.89")
    assert _importe("no consta") is None
    assert _importe(None) is None


def test_la_hoja_de_presupuestos_ordena_coincidencias_diferencias_y_sin_importe():
    from app.presupuestos_adif import (
        COINCIDE, DIFIERE, SIN_IMPORTE_LEIDO, FilaPresupuestoAdif, ordenar,
    )

    def fila(codigo, resultado, diferencia):
        return FilaPresupuestoAdif(
            codigo_expediente=codigo, titulo=None, presupuesto_adif=Decimal("1"),
            importe_documentos=None, diferencia=diferencia, diferencia_relativa=None,
            resultado=resultado,
        )

    filas = [
        fila("c", SIN_IMPORTE_LEIDO, None),
        fila("b", DIFIERE, Decimal("-900000")),
        fila("a", DIFIERE, Decimal("10")),
        fila("d", COINCIDE, Decimal("0")),
    ]
    assert [f.codigo_expediente for f in ordenar(filas)] == ["d", "b", "a", "c"]


# --------------------------------------------------------------------------
# 3 — El catálogo antiguo de ADIF
# --------------------------------------------------------------------------


def test_la_descripcion_normalizada_ignora_acentos_mayusculas_y_puntuacion():
    assert clave_descripcion("BRIDA de unión, D-60 mm.") == clave_descripcion("brida de  union d 60 mm")
    assert clave_descripcion("   ") is None
    assert clave_descripcion(None) is None


def test_el_catalogo_antiguo_se_lee_aunque_el_encabezado_no_este_en_la_primera_fila(tmp_path):
    ruta = _libro(
        [
            ("CATÁLOGO DE MATERIALES ADIF", None, None),
            (None, None, None),
            ("Material", "Denominación", "Precio unitario"),
            ("650000030", "Aire acondicionado 1x1", "3.100,25 €"),
            (None, "Brida de unión", 12.5),
        ],
        ruta=tmp_path / "catalogo.xlsx",
    )
    materiales, columnas, hojas = leer_catalogo_antiguo(ruta)
    assert columnas == {"matricula": 0, "descripcion": 1, "precio": 2}
    assert hojas == ["Hoja1"]
    assert [m.matricula for m in materiales] == ["650000030", None]
    assert materiales[0].precio == Decimal("3100.25")
    assert materiales[0].fila == 4


def test_un_catalogo_sin_matricula_ni_descripcion_no_se_lee(tmp_path):
    ruta = _libro([("Año", "Total"), (2024, 10)], ruta=tmp_path / "otro.xlsx")
    materiales, _columnas, hojas = leer_catalogo_antiguo(ruta)
    assert materiales == [] and hojas == []


def test_una_columna_de_total_no_pasa_por_precio_unitario(tmp_path):
    ruta = _libro(
        [("Material", "Descripción", "Importe total del pedido"), ("650000030", "Aire", 100.0)],
        ruta=tmp_path / "catalogo.xlsx",
    )
    _materiales, columnas, _hojas = leer_catalogo_antiguo(ruta)
    assert columnas["precio"] is None


def test_el_informe_sale_con_sus_tres_hojas_y_su_resumen(tmp_path):
    informe = InformeCatalogoAntiguo(
        resumen=ResumenCatalogoAntiguo(configurado=True, formato_reconocido=True),
        solo_en_el_suyo=[MaterialCatalogoAntiguo("111111111", "Suyo", Decimal("1"), "H", 2)],
        solo_en_el_nuestro=[MaterialNuestro("222222222", "Nuestro", Decimal("2"), "6.24/28510.0088", "1")],
    )
    contenido = generar_informe_catalogo_antiguo(informe)
    ruta = tmp_path / "informe.xlsx"
    ruta.write_bytes(contenido)
    libro = openpyxl.load_workbook(ruta)
    assert libro.sheetnames == [
        "Solo en el catálogo de ADIF", "Solo en el nuestro", "Diferencias de precio", "Resumen",
    ]
    assert libro["Solo en el catálogo de ADIF"].cell(row=2, column=1).value == "111111111"

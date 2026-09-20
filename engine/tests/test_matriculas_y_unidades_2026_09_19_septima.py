"""Sesión 2026-09-19 (séptima parte): las matrículas que no figuran en el
maestro de materiales de ADIF y las unidades raras del Excel.

Bloque 1 — las matrículas de 8 dígitos. El documento las imprime de verdad
(comprobado contra la imagen de los 20 expedientes afectados), así que se
quedan literales y la fila lleva su motivo. **Nunca se completan con un
dígito para que casen con el maestro**: `64571017` casa a la vez con
`645710170` (una palomilla) y con `645710175` (las antenas).

Bloque 3 — "Ml" es metro lineal y se unifica con "m".
"""
import io
from decimal import Decimal

import openpyxl
import pytest

from app.catalogo import (
    MOTIVO_MATRICULA_8_FUERA_DEL_MAESTRO,
    MOTIVO_MATRICULA_FUERA_DEL_MAESTRO,
    construir_linea_catalogo,
    guardar_lineas_catalogo,
    matriculas_fuera_del_maestro,
)
from app.exportacion import _sale_en_materiales, generar_excel_catalogo
from app.extraccion.unidad_medida import (
    es_unidad_conocida,
    limpiar_unidad,
    normalizar_unidad,
)
from app.models import Expediente, LineaCatalogo, Lote, MaestroMaterial

_MAPEO = {
    "codigo_precio": 0,
    "matricula": 1,
    "descripcion": 2,
    "unidad_medida": 3,
    "cantidad": 4,
    "precio_unitario": 5,
}


def _lote(db_session, codigo="6.16/28510.9042"):
    expediente = Expediente(codigo_expediente=codigo)
    db_session.add(expediente)
    db_session.commit()
    lote = Lote(expediente_id=expediente.id, identificador_lote="1")
    db_session.add(lote)
    db_session.commit()
    return lote


def _maestro(db_session, *matriculas):
    for matricula in matriculas:
        db_session.add(
            MaestroMaterial(matricula=matricula, descripcion="X", unidad_medida="UN")
        )
    db_session.commit()


def _guardar(db_session, lote, filas):
    lineas = []
    for orden, fila in enumerate(filas):
        linea = construir_linea_catalogo(
            fila,
            _MAPEO,
            pagina=11,
            documento_origen_id=None,
            expediente_id=lote.expediente_id,
            baja_lote=None,
            orden_aparicion=orden,
        )
        assert linea is not None
        lineas.append(linea)
    resultado = guardar_lineas_catalogo(db_session, lote.id, lineas)
    db_session.commit()
    return resultado


# --------------------------------------------------------------------------
# Bloque 1 — la matrícula de 8 dígitos que no está en el maestro
# --------------------------------------------------------------------------


def test_la_matricula_de_8_que_no_esta_en_el_maestro_lleva_su_motivo(db_session):
    # `6.18/28510.0066`: su pliego imprime "59020019" y el maestro de hoy no
    # lo tiene. La celda NO se toca -- la matrícula sale literal.
    _maestro(db_session, "590200190", "697500900")
    lote = _lote(db_session)

    _guardar(db_session, lote, [["P-001", "59020019", "FUSIBLE AT - CA", "UN", "1", "116,15"]])

    guardada = db_session.query(LineaCatalogo).filter_by(lote_id=lote.id).one()
    assert guardada.matricula == "59020019"
    assert MOTIVO_MATRICULA_8_FUERA_DEL_MAESTRO in guardada.motivo_revision


def test_la_fila_con_ese_motivo_sigue_saliendo_en_el_entregable():
    # El motivo explica, no excluye: en almacenes tiene que verse la fila con
    # su matrícula, igual que cualquier otra.
    assert _sale_en_materiales(
        MOTIVO_MATRICULA_8_FUERA_DEL_MAESTRO, tiene_lote=True, incluir_pendientes_sin_lote=False
    )


def test_la_matricula_de_8_que_si_esta_en_el_maestro_no_lleva_motivo(db_session):
    # Las 440 matrículas cortas del maestro (categorías genéricas de SAP) no
    # son ninguna anomalía: si figura, no hay nada que explicar.
    _maestro(db_session, "59020019")
    lote = _lote(db_session)

    _guardar(db_session, lote, [["P-001", "59020019", "FUSIBLE AT - CA", "UN", "1", "116,15"]])

    guardada = db_session.query(LineaCatalogo).filter_by(lote_id=lote.id).one()
    assert not guardada.motivo_revision


def test_la_matricula_de_9_lleva_su_propio_motivo(db_session):
    # Bloque 2 del encargo: las de 9 dígitos que el documento imprime así y el
    # maestro no recoge llevan su motivo, DISTINTO del de las de 8 -- una es
    # un formato que ADIF ya no usa, la otra es el formato de hoy al que le
    # falta la fila del maestro.
    _maestro(db_session, "111111111")
    lote = _lote(db_session)

    _guardar(db_session, lote, [["P-001", "645710175", "ANTENAS DE ARQUEO", "UN", "10", "70,00"]])

    guardada = db_session.query(LineaCatalogo).filter_by(lote_id=lote.id).one()
    assert guardada.matricula == "645710175"
    assert MOTIVO_MATRICULA_FUERA_DEL_MAESTRO in guardada.motivo_revision
    assert MOTIVO_MATRICULA_8_FUERA_DEL_MAESTRO not in guardada.motivo_revision


def test_la_matricula_de_9_que_si_esta_en_el_maestro_no_lleva_motivo(db_session):
    _maestro(db_session, "645710175")
    lote = _lote(db_session)

    _guardar(db_session, lote, [["P-001", "645710175", "ANTENAS DE ARQUEO", "UN", "10", "70,00"]])

    guardada = db_session.query(LineaCatalogo).filter_by(lote_id=lote.id).one()
    assert not guardada.motivo_revision


def test_la_fila_de_9_con_ese_motivo_sigue_saliendo_en_el_entregable():
    assert _sale_en_materiales(
        MOTIVO_MATRICULA_FUERA_DEL_MAESTRO, tiene_lote=True, incluir_pendientes_sin_lote=False
    )


def test_sin_maestro_cargado_no_se_escribe_el_motivo(db_session):
    # Sin listado contra el que comprobar no se puede afirmar que una
    # matrícula no figure en él: marcar las 390 filas por el hueco de una
    # fuente de entrada sería escribir un motivo falso.
    lote = _lote(db_session)

    _guardar(db_session, lote, [["P-001", "59020019", "FUSIBLE AT - CA", "UN", "1", "116,15"]])

    guardada = db_session.query(LineaCatalogo).filter_by(lote_id=lote.id).one()
    assert not guardada.motivo_revision


def test_reprocesar_no_repite_el_motivo(db_session):
    # Invariante 9: el ciclo de mantenimiento reextrae los mismos documentos
    # cada semana y el texto del motivo no puede ir creciendo.
    _maestro(db_session, "590200190")
    lote = _lote(db_session)
    filas = [["P-001", "59020019", "FUSIBLE AT - CA", "UN", "1", "116,15"]]

    _guardar(db_session, lote, filas)
    _guardar(db_session, lote, filas)

    guardada = db_session.query(LineaCatalogo).filter_by(lote_id=lote.id).one()
    assert guardada.motivo_revision.count(MOTIVO_MATRICULA_8_FUERA_DEL_MAESTRO) == 1


def test_el_motivo_desaparece_si_la_matricula_pasa_a_estar_en_el_maestro(db_session):
    # El día que ADIF mande un maestro que ya las traiga, el motivo se retira
    # solo: `motivo_revision` se recalcula entero en cada pasada.
    _maestro(db_session, "590200190")
    lote = _lote(db_session)
    filas = [["P-001", "59020019", "FUSIBLE AT - CA", "UN", "1", "116,15"]]
    _guardar(db_session, lote, filas)

    _maestro(db_session, "59020019")
    _guardar(db_session, lote, filas)

    guardada = db_session.query(LineaCatalogo).filter_by(lote_id=lote.id).one()
    assert not guardada.motivo_revision


def test_matriculas_fuera_del_maestro_devuelve_las_que_faltan(db_session):
    _maestro(db_session, "590200190")

    fuera = matriculas_fuera_del_maestro(
        db_session, ["59020019", "590200190", "645710175", None, "1234"]
    )

    assert fuera == frozenset({"59020019", "645710175", "1234"})


# --------------------------------------------------------------------------
# Bloque 3 — "Ml" es metro lineal y se unifica con "m"
# --------------------------------------------------------------------------


@pytest.mark.parametrize("literal", ["Ml", "ML", "ml", "m.l.", "M.L."])
def test_metro_lineal_se_unifica_con_metro(literal):
    # `6.17/28510.0007` p.20: obra civil medida a lo largo (lámina geotextil,
    # muro de contención, zona de paso), en un cuadro que usa m3 y m2 en las
    # filas de al lado. No es mililitro.
    assert es_unidad_conocida(literal)
    assert normalizar_unidad(literal) == "m"


def test_el_literal_del_documento_no_se_pierde(db_session):
    # `unidad_medida_original` guarda lo que imprime el documento; la columna
    # del Excel lleva la forma única. Así la unificación es reversible.
    lote = _lote(db_session)

    _guardar(db_session, lote, [["P-006", None, "Colocación de lámina geotextil.", "Ml", "800", "1,10"]])

    guardada = db_session.query(LineaCatalogo).filter_by(lote_id=lote.id).one()
    assert guardada.unidad_medida == "m"
    assert guardada.unidad_medida_original == "Ml"


def test_metro_por_litro_no_se_toca():
    # La unificación es de "ml" como una sola pieza, no de una compuesta que
    # divida metros entre litros.
    assert normalizar_unidad("m/l") == "m/l"


def test_transporte_sigue_siendo_la_unidad_que_imprime_el_documento():
    # `6.21/28510.0108`-`0112`: el cuadro escribe "€/transporte" en la misma
    # columna en la que escribe "€/Ton*km" para las filas de tonelada-
    # kilómetro. Es el denominador del precio, no un concepto colado de otra
    # columna: se queda como está mientras el cliente no diga otra cosa.
    assert es_unidad_conocida("€/transporte")
    assert normalizar_unidad(limpiar_unidad("€/transporte")) == "transporte"


# --------------------------------------------------------------------------
# El Excel: los dos motivos nuevos, contados en el Resumen
# --------------------------------------------------------------------------


def _linea(lote, clave, matricula, motivo):
    return LineaCatalogo(
        expediente_id=lote.expediente_id, lote_id=lote.id, clave_linea=clave,
        orden_aparicion=0, matricula=matricula, descripcion="FUSIBLE AT - CA",
        precio_unitario=Decimal("116.15"), motivo_revision=motivo,
    )


def _resumen(db_session):
    libro = openpyxl.load_workbook(io.BytesIO(generar_excel_catalogo(db_session)))
    return libro, [
        (str(f[0].value), f[1].value if len(f) > 1 else None)
        for f in libro["Resumen"].iter_rows()
        if f[0].value is not None
    ]


def test_el_resumen_cuenta_las_dos_familias_de_matricula_por_separado(db_session):
    lote = _lote(db_session)
    db_session.add_all([
        _linea(lote, "A", "59020019", MOTIVO_MATRICULA_8_FUERA_DEL_MAESTRO),
        _linea(lote, "B", "645710175", MOTIVO_MATRICULA_FUERA_DEL_MAESTRO),
        _linea(lote, "C", "645710176", MOTIVO_MATRICULA_FUERA_DEL_MAESTRO),
    ])
    db_session.commit()

    libro, filas = _resumen(db_session)

    ocho = [v for t, v in filas if "formato antiguo de 8 dígitos" in t]
    nueve = [v for t, v in filas if "(9 dígitos) no figura en el maestro" in t]
    assert ocho == [1]
    assert nueve == [2]
    # Y las tres líneas salen en "Materiales" con su matrícula: el motivo
    # explica, no excluye.
    assert libro["Materiales"].max_row == 4


def test_el_resumen_explica_por_que_no_se_completan(db_session):
    lote = _lote(db_session)
    db_session.add(_linea(lote, "A", "59020019", MOTIVO_MATRICULA_8_FUERA_DEL_MAESTRO))
    db_session.commit()

    _, filas = _resumen(db_session)

    texto = "\n".join(t for t, _ in filas)
    assert "no se completan ni se corrigen" in texto
    assert "64571017" in texto


def test_sin_ninguna_linea_asi_el_resumen_no_escribe_la_nota(db_session):
    lote = _lote(db_session)
    db_session.add(_linea(lote, "A", "601200010", None))
    db_session.commit()

    _, filas = _resumen(db_session)

    texto = "\n".join(t for t, _ in filas)
    assert "no se completan ni se corrigen" not in texto
    assert [v for t, v in filas if "formato antiguo de 8 dígitos" in t] == [0]


# --------------------------------------------------------------------------
# Las 11 filas que se quedaban sin motivo (cierre de la sesión, 2026-09-20)
#
# El motivo se calculaba sobre `datos`, antes de escribir la fila. Dos
# caminos se le escapaban, y entre los dos dejaban 11 filas del entregable
# fuera del maestro y sin explicación. Ahora se calcula sobre la fila ya
# escrita, que es lo que el Excel muestra.
# --------------------------------------------------------------------------


def _heredada(lote, matricula, motivo=None):
    """Una línea como la que escribe `app.extraccion.herencia_matriz`: sin
    `motivo_revision` en el diccionario, porque ahí la ausencia significa "no
    se ha evaluado" y no debe pisar lo que la fila tuviera."""
    datos = {
        "clave_linea": f"P-{matricula}",
        "expediente_id": lote.expediente_id,
        "orden_aparicion": 0,
        "matricula": matricula,
        "descripcion": "CABLE CONEXIÓN FLEXIBLE 120 MM2",
        "precio_unitario": Decimal("12.34"),
        "heredado_de_matriz": True,
    }
    if motivo is not None:
        datos["motivo_revision"] = motivo
    return datos


def test_la_linea_heredada_de_un_acuerdo_marco_tambien_lleva_el_motivo(db_session):
    # Nueve de las once: `642950151` en ocho expedientes y `642190440` en
    # `6.20/28510.0040`. Son líneas que el pedido no lee de su propio
    # documento, sino que copia de su matriz, y salen en "Materiales" con su
    # matrícula como cualquier otra -- así que necesitan la misma explicación.
    _maestro(db_session, "642950150")
    lote = _lote(db_session)

    guardar_lineas_catalogo(db_session, lote.id, [_heredada(lote, "642950151")])
    db_session.commit()

    guardada = db_session.query(LineaCatalogo).filter_by(lote_id=lote.id).one()
    assert guardada.heredado_de_matriz
    assert guardada.motivo_revision == MOTIVO_MATRICULA_FUERA_DEL_MAESTRO


def test_en_la_heredada_el_motivo_tambien_se_retira_solo(db_session):
    # La condición que más importa de las cuatro: el motivo tiene que irse
    # solo el día que ADIF mande un maestro más completo. En estas líneas
    # `motivo_revision` no se recalcula de cero en cada pasada, así que hay
    # que quitarlo explícitamente antes de volver a ponerlo.
    _maestro(db_session, "642950150")
    lote = _lote(db_session)
    guardar_lineas_catalogo(db_session, lote.id, [_heredada(lote, "642950151")])
    db_session.commit()

    _maestro(db_session, "642950151")
    guardar_lineas_catalogo(db_session, lote.id, [_heredada(lote, "642950151")])
    db_session.commit()

    guardada = db_session.query(LineaCatalogo).filter_by(lote_id=lote.id).one()
    assert not guardada.motivo_revision


def test_en_la_heredada_el_motivo_no_pisa_el_que_la_fila_ya_tenia(db_session):
    # Quitar y poner el motivo del maestro no puede llevarse por delante el
    # resto: la herencia deja `motivo_revision` fuera de `datos` justamente
    # para no tocarlo.
    _maestro(db_session, "642950150")
    lote = _lote(db_session)
    guardar_lineas_catalogo(
        db_session, lote.id, [_heredada(lote, "642950151", motivo="cantidad es 0")]
    )
    db_session.commit()

    guardar_lineas_catalogo(db_session, lote.id, [_heredada(lote, "642950151")])
    db_session.commit()

    guardada = db_session.query(LineaCatalogo).filter_by(lote_id=lote.id).one()
    assert guardada.motivo_revision == f"cantidad es 0; {MOTIVO_MATRICULA_FUERA_DEL_MAESTRO}"


def test_la_matricula_que_la_pasada_no_reconoce_pero_la_fila_conserva(db_session):
    # Las otras dos: `611050007` en `6.22/28510.0094` y `619050222` en
    # `6.22/28510.0125`. Una pasada que no reconoce la celda deja `matricula`
    # a `None`, y un `None` corriente no pisa un valor ya conocido -- así que
    # la fila se entrega con la matrícula que escribió otra pasada. El motivo
    # tiene que hablar de ESA, la que el Excel enseña.
    _maestro(db_session, "611050075")
    lote = _lote(db_session)
    _guardar(db_session, lote, [["P-001", "611050007", "CARRIL UIC-54", "UN", "1", "10,00"]])

    # Segunda pasada: la celda viene ilegible y la matrícula se descarta.
    _guardar(db_session, lote, [["P-001", "6,111E+09", "CARRIL UIC-54", "UN", "1", "10,00"]])

    guardada = db_session.query(LineaCatalogo).filter_by(lote_id=lote.id).one()
    assert guardada.matricula == "611050007"
    assert MOTIVO_MATRICULA_FUERA_DEL_MAESTRO in guardada.motivo_revision


def test_sin_maestro_cargado_la_heredada_conserva_el_motivo_que_ya_tenia(db_session):
    # Simétrico de "sin maestro no se escribe nada": sin listado contra el
    # que comprobar, esta función ni pone ni quita.
    #
    # Ojo con el alcance: esto vale para la línea heredada, que es la que no
    # recalcula su `motivo_revision`. En el camino normal el motivo SÍ
    # desaparece si el maestro deja de estar cargado, y no por esta función
    # sino porque `_construir_campos` reconstruye `motivo_revision` entero en
    # cada pasada -- comportamiento de siempre y coherente con la misma
    # regla: sin el listado no se puede afirmar nada sobre la matrícula.
    _maestro(db_session, "642950150")
    lote = _lote(db_session)
    guardar_lineas_catalogo(db_session, lote.id, [_heredada(lote, "642950151")])
    db_session.commit()

    db_session.query(MaestroMaterial).delete()
    db_session.commit()
    guardar_lineas_catalogo(db_session, lote.id, [_heredada(lote, "642950151")])
    db_session.commit()

    guardada = db_session.query(LineaCatalogo).filter_by(lote_id=lote.id).one()
    assert guardada.motivo_revision == MOTIVO_MATRICULA_FUERA_DEL_MAESTRO


def test_si_la_matricula_cambia_de_longitud_el_motivo_cambia_con_ella(db_session):
    # Lo que hizo esta misma sesión 34 veces: corregir una lectura que perdía
    # un dígito. La fila pasa de "formato antiguo de 8 dígitos" al motivo de
    # 9, sin quedarse los dos pegados.
    _maestro(db_session, "111111111")
    lote = _lote(db_session)
    _guardar(db_session, lote, [["P-001", "64571017", "ANTENAS DE ARQUEO", "UN", "1", "70,00"]])
    guardada = db_session.query(LineaCatalogo).filter_by(lote_id=lote.id).one()
    assert MOTIVO_MATRICULA_8_FUERA_DEL_MAESTRO in guardada.motivo_revision

    _guardar(db_session, lote, [["P-001", "645710175", "ANTENAS DE ARQUEO", "UN", "1", "70,00"]])

    guardada = db_session.query(LineaCatalogo).filter_by(lote_id=lote.id).one()
    assert guardada.motivo_revision == MOTIVO_MATRICULA_FUERA_DEL_MAESTRO

"""Pruebas de la API de catálogo, revisión y documentos (CONTEXTO.md, encargo
de esta sesión, puntos 1 a 4): sobre `TestClient` con SQLite en memoria, sin
Postgres levantado — mismo patrón que `db_session` de conftest.py."""
import io
from decimal import Decimal

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import models
from app.catalogo import MOTIVO_VALOR_DE_OTRO_LOTE
from app.config import settings
from app.db import Base, get_db
from app.interfaces.document_storage import LocalDiskStorage
from app.main import app
from app.models import (
    Documento,
    DocumentoExpediente,
    EstadoExpediente,
    EstadoRevisionLinea,
    Expediente,
    LineaCatalogo,
    Lote,
    TipoDocumento,
)

_TABLAS = [
    models.Expediente.__table__,
    models.Lote.__table__,
    models.Documento.__table__,
    models.DocumentoExpediente.__table__,
    models.LineaCatalogo.__table__,
    models.TrazaOrigen.__table__,
    models.MapeoCabeceraCache.__table__,
    models.TrabajoCola.__table__,
]


@pytest.fixture()
def db_session():
    # TestClient corre el ASGI app en un hilo de threadpool distinto del
    # hilo del test (starlette.concurrency.run_in_threadpool); una conexión
    # SQLite normal no se puede compartir entre hilos, así que hace falta
    # una única conexión fija (StaticPool + check_same_thread=False) en vez
    # de una nueva conexión ":memory:" (vacía) por hilo.
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine, tables=_TABLAS)
    sesion = sessionmaker(bind=engine)()
    try:
        yield sesion
    finally:
        sesion.close()


@pytest.fixture()
def cliente(db_session):
    def _get_db_prueba():
        yield db_session

    app.dependency_overrides[get_db] = _get_db_prueba
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.clear()


def _sembrar_catalogo(db_session, *, estado=EstadoExpediente.completado):
    expediente = Expediente(
        codigo_expediente="6.24/28510.0008",
        codigo_matriz=None,
        nombre_proyecto="SUMINISTRO DE GUANTES CONTRA RIESGO ELECTRICO.",
        estado=estado,
        error="no se pudo determinar la baja del lote" if estado == EstadoExpediente.pendiente_revision else None,
    )
    db_session.add(expediente)
    db_session.commit()

    lote = Lote(expediente_id=expediente.id, identificador_lote="1", baja_lote=Decimal("0.5400"))
    db_session.add(lote)
    db_session.commit()

    documento = Documento(
        tipo_documento=TipoDocumento.anejo,
        hash="hash-1",
        ruta_almacenamiento="6.24_28510.0008_ANEJO_1.pdf",
    )
    db_session.add(documento)
    db_session.commit()
    db_session.add(DocumentoExpediente(
        documento_id=documento.id,
        expediente_id=expediente.id,
        nombre_archivo="6.24_28510.0008_ANEJO_1.pdf",
    ))
    db_session.commit()

    linea = LineaCatalogo(
        lote_id=lote.id,
        expediente_id=expediente.id,
        clave_linea="P-001",
        orden_aparicion=0,
        codigo_precio="P-001",
        matricula="697500900",
        descripcion="GUANTE CONTRA RIESGO ELECTRICO",
        codigo_material="GUANTE",
        cantidad=Decimal("30"),
        precio_unitario=Decimal("24.00"),
        unidad_medida="UN",
        baja_lote=Decimal("0.5400"),
        precio_adjudicado=Decimal("11.0400"),
        documento_origen_id=documento.id,
        pagina=11,
        fragmento="P-001 | GUANTE...",
    )
    db_session.add(linea)
    db_session.commit()
    return expediente, lote, documento, linea


def test_catalogo_lista_lineas_con_datos_de_expediente(cliente, db_session):
    _sembrar_catalogo(db_session)

    resp = cliente.get("/catalogo")

    assert resp.status_code == 200
    datos = resp.json()
    assert datos["total"] == 1
    fila = datos["lineas"][0]
    assert fila["codigo_expediente"] == "6.24/28510.0008"
    assert fila["matricula"] == "697500900"
    assert fila["precio_adjudicado"] == "11.0400"
    assert fila["documento_origen_nombre"] == "6.24_28510.0008_ANEJO_1.pdf"
    assert fila["pagina"] == 11


def test_catalogo_busqueda_por_matricula_a_traves_de_expedientes(cliente, db_session):
    _sembrar_catalogo(db_session)
    otro = Expediente(codigo_expediente="6.24/28510.9999", estado=EstadoExpediente.completado)
    db_session.add(otro)
    db_session.commit()
    otro_lote = Lote(expediente_id=otro.id, identificador_lote="1")
    db_session.add(otro_lote)
    db_session.commit()
    db_session.add(LineaCatalogo(
        lote_id=otro_lote.id, expediente_id=otro.id, clave_linea="P-099", orden_aparicion=0,
        matricula="697500900", descripcion="GUANTE OTRO EXPEDIENTE", precio_unitario=Decimal("25.00"),
    ))
    db_session.commit()

    resp = cliente.get("/catalogo", params={"matricula": "697500900"})

    assert resp.status_code == 200
    datos = resp.json()
    assert datos["total"] == 2
    expedientes = {f["codigo_expediente"] for f in datos["lineas"]}
    assert expedientes == {"6.24/28510.0008", "6.24/28510.9999"}


def test_catalogo_filtra_por_expediente_sin_afectar_a_otros(cliente, db_session):
    _sembrar_catalogo(db_session)

    resp = cliente.get("/catalogo", params={"expediente": "6.24/28510.9999"})

    assert resp.json()["total"] == 0


def _sembrar_linea_escasa(db_session, *, codigo_expediente):
    # Aviso del cliente (sesión 2026-09-07): una línea sin código de precio,
    # sin matrícula, sin cantidad y sin unidad -- el caso real de
    # `6.20/28510.0054`/`6.20/28510.0136`, tablas de origen genuinamente sin
    # esas columnas. `codigo_expediente` se elige para que ordene ANTES que
    # `6.24/28510.0008` (la línea completa que ya siembra `_sembrar_catalogo`)
    # en orden alfabético, y así distinguir de verdad los dos criterios.
    expediente = Expediente(codigo_expediente=codigo_expediente, estado=EstadoExpediente.completado)
    db_session.add(expediente)
    db_session.commit()
    lote = Lote(expediente_id=expediente.id, identificador_lote="1")
    db_session.add(lote)
    db_session.commit()
    linea = LineaCatalogo(
        lote_id=lote.id, expediente_id=expediente.id, clave_linea="escasa",
        orden_aparicion=0, descripcion="LINEA SIN CASI NADA", precio_unitario=Decimal("1.00"),
    )
    db_session.add(linea)
    db_session.commit()
    return expediente, linea


def test_catalogo_orden_por_defecto_pone_primero_las_lineas_completas(cliente, db_session):
    # Encargo de esta sesión: antes de este arreglo, el orden alfabético
    # ponía siempre primero a expedientes con tablas de origen escasas
    # (código de expediente que ordena antes), aunque el resto del catálogo
    # estuviera lleno -- la primera pantalla sin filtrar, la que ve el
    # cliente, daba la impresión contraria a la realidad.
    _sembrar_linea_escasa(db_session, codigo_expediente="6.20/28510.0001")
    *_, completa = _sembrar_catalogo(db_session)  # "6.24/28510.0008", todos los campos rellenos

    resp = cliente.get("/catalogo")

    assert resp.status_code == 200
    lineas = resp.json()["lineas"]
    assert lineas[0]["id"] == completa.id
    assert lineas[-1]["descripcion"] == "LINEA SIN CASI NADA"


def test_catalogo_orden_alfabetico_sigue_disponible(cliente, db_session):
    # El criterio de siempre se deja disponible sin más que pedirlo --
    # mismo resultado que antes de este arreglo.
    escasa, _ = _sembrar_linea_escasa(db_session, codigo_expediente="6.20/28510.0001")
    _sembrar_catalogo(db_session)  # "6.24/28510.0008"

    resp = cliente.get("/catalogo", params={"orden": "alfabetico"})

    assert resp.status_code == 200
    lineas = resp.json()["lineas"]
    assert lineas[0]["codigo_expediente"] == escasa.codigo_expediente


def test_catalogo_orden_invalido_rechazado(cliente, db_session):
    resp = cliente.get("/catalogo", params={"orden": "aleatorio"})

    assert resp.status_code == 422


def test_exportar_catalogo_genera_xlsx_con_columnas_del_formato_esperado(cliente, db_session, tmp_path, monkeypatch):
    # Este test comprueba el caso "sin cruce" (fila[0]/fila[1] vacíos) --
    # tiene que valer sea cual sea el entorno local, incluso con un
    # docker-compose.override.yml real montando un Excel de códigos de
    # verdad (CODIGOS_PROYECTO_PATH no vacío fuera de este proceso de test).
    monkeypatch.setattr(settings, "codigos_proyecto_path", None)
    _sembrar_catalogo(db_session)

    resp = cliente.get("/catalogo/exportar.xlsx")

    assert resp.status_code == 200
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    ruta = tmp_path / "salida.xlsx"
    ruta.write_bytes(resp.content)
    libro = openpyxl.load_workbook(ruta)
    hoja = libro.active
    cabecera = [c.value for c in next(hoja.iter_rows(min_row=1, max_row=1))]
    # CONTEXTO.md, encargo de esta sesión, punto 4: mismas columnas y mismo
    # orden que Ejemplo/Output/.
    assert cabecera == [
        "Código interno", "Código de expediente", "Código matriz", "Título expediente",
        "Matrícula del material", "Descripción del material", "Código del material",
        "Cantidad", "Precio unitario", "Lote",
        "Precio adjudicado", "Baja del lote", "Unidad de medida",
        "Estado del contrato (SAP)",
        "Nº de expediente (documento)", "Objeto del contrato (documento)",
        "Motivo de las celdas vacías",
        "Comentarios",
    ]
    fila = [c.value for c in next(hoja.iter_rows(min_row=2, max_row=2))]
    assert fila[3] == "SUMINISTRO DE GUANTES CONTRA RIESGO ELECTRICO."
    assert fila[4] == "697500900"
    assert fila[5] == "GUANTE CONTRA RIESGO ELECTRICO"
    assert fila[8] == 24.0
    # Sin cruce con el Excel de códigos configurado en este test: código
    # interno y código de proyecto se dejan vacíos, no inventados
    # (CONTEXTO.md sección 7) -- un espacio, no una celda en blanco del todo
    # (bloque 2, segunda tanda de cambios del cliente, sesión 2026-09-09:
    # solo en columnas de texto, para que Excel no desborde encima el texto
    # de la celda anterior).
    assert fila[0] == " "
    assert fila[1] == " "
    # Encargo de una sesión anterior, punto 3: precio adjudicado y baja del
    # lote, sin desplazar las once columnas de siempre.
    assert fila[10] == 11.04
    assert fila[11] == 0.54
    # Aviso del cliente (sesión 2026-09-07): unidad de medida, al final de
    # todo, después de las dos columnas ya añadidas.
    assert fila[12] == "UN"
    # Bloque 1, sesión de comparación documento-vs-listado interno: el
    # número de expediente y el objeto/título del documento SIEMPRE se
    # rellenan, aunque el expediente no haya cruzado con el Excel de
    # códigos (a diferencia de fila[0]/fila[1] de arriba, vacíos en este
    # mismo test por la misma razón).
    assert fila[14] == "6.24/28510.0008"
    assert fila[15] == "SUMINISTRO DE GUANTES CONTRA RIESGO ELECTRICO."
    # Sesión 2026-09-14 (continuación): la línea sembrada trae todos sus
    # datos, así que no hay ninguna celda vacía que explicar.
    assert fila[16] == " "
    # Segunda tanda de cambios del cliente (bloque 1, sesión 2026-09-09):
    # "Comentarios" se mueve al final de todas las columnas.
    assert fila[17] == " "


def test_exportar_catalogo_explica_el_hueco_de_un_valor_de_otro_lote(cliente, db_session, tmp_path, monkeypatch):
    """Encargo del cliente (sesión 2026-09-14, continuación): una cantidad
    vacía porque el documento da una distinta para cada lote no puede leerse
    igual que una que el documento no trae. La celda sigue vacía (y la
    columna numérica), el motivo va en su columna con el criterio de los
    tres motivos, y la hoja Resumen lo cuenta y lo explica."""
    monkeypatch.setattr(settings, "codigos_proyecto_path", None)
    _expediente, _lote, _documento, linea = _sembrar_catalogo(db_session)
    linea.cantidad = None
    linea.matricula = None
    linea.motivo_revision = (
        f"{MOTIVO_VALOR_DE_OTRO_LOTE}, con cantidades distintos: probablemente cuadros de lotes distintos "
        "adjuntos al mismo documento -- no se puede atribuir un valor a este lote con certeza, se deja vacío "
        "en vez de quedarse con el último visto"
    )
    db_session.commit()

    resp = cliente.get("/catalogo/exportar.xlsx")

    ruta = tmp_path / "salida.xlsx"
    ruta.write_bytes(resp.content)
    libro = openpyxl.load_workbook(ruta)
    fila = [c.value for c in next(libro["Materiales"].iter_rows(min_row=2, max_row=2))]
    assert fila[7] is None  # Cantidad: vacía, no un texto en una columna numérica
    assert fila[16] == (
        "Matrícula del material: no consta; "
        "Cantidad: pendiente (el documento da una cantidad distinta para cada lote y falta saber cuál es la de este)"
    )
    filas_resumen = [[c.value for c in f] for f in libro["Resumen"].iter_rows()]
    assert any(
        (f[0] or "").startswith("Líneas del catálogo con Cantidad o Precio unitario pendiente") and f[1] == 1
        for f in filas_resumen
    )
    assert any((f[0] or "").startswith("Pendiente: ") for f in filas_resumen)

    # La web recibe el mismo motivo.
    datos = cliente.get("/catalogo").json()["lineas"][0]
    assert datos["celdas_vacias"]["cantidad"]["motivo"] == "pendiente"
    assert datos["celdas_vacias"]["matricula"] == {"motivo": "no-consta", "detalle": None}
    assert "precio_unitario" not in datos["celdas_vacias"]


def _agregar_linea_huerfana(db_session, expediente, documento, *, motivo):
    linea = LineaCatalogo(
        lote_id=None,
        expediente_id=expediente.id,
        clave_linea="P-002@p5y100",
        orden_aparicion=1,
        codigo_precio="P-002",
        descripcion="TORNILLO M8",
        precio_unitario=Decimal("3.00"),
        documento_origen_id=documento.id,
        pagina=5,
        fragmento="P-002 | TORNILLO M8...",
        motivo_revision=motivo,
    )
    db_session.add(linea)
    db_session.commit()
    return linea


def test_exportar_catalogo_excluye_huerfanas_por_defecto_y_las_resume(cliente, db_session, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "codigos_proyecto_path", None)
    expediente, _lote, documento, _linea = _sembrar_catalogo(db_session)
    _agregar_linea_huerfana(
        db_session, expediente, documento,
        motivo="banda vacía: posible continuación de tabla partida entre páginas, sin inferir",
    )

    resp = cliente.get("/catalogo/exportar.xlsx")

    ruta = tmp_path / "salida.xlsx"
    ruta.write_bytes(resp.content)
    libro = openpyxl.load_workbook(ruta)
    materiales = libro["Materiales"]
    # Encargo de esta sesión: la huérfana (sin lote asignado con seguridad)
    # no sale en "Materiales" por defecto -- solo la línea con lote.
    assert materiales.max_row == 2
    descripciones = [c.value for c in materiales["F"]][1:]
    assert "TORNILLO M8" not in descripciones

    resumen = libro["Resumen"]
    filas_resumen = [[c.value for c in fila] for fila in resumen.iter_rows()]
    assert ["Líneas en este catálogo", 1, None] in filas_resumen
    assert ["Líneas pendientes de revisión (no incluidas arriba)", 1, None] in filas_resumen
    # Encargo de esta sesión, punto 4: el motivo se explica en lenguaje llano
    # (qué ha pasado / qué haría falta), no con la etiqueta técnica interna.
    assert any(
        fila[1] == 1 and "continuar de una página a la siguiente" in (fila[0] or "")
        for fila in filas_resumen
    )
    assert any("Código del material" in (fila[0] or "") for fila in filas_resumen)


def test_exportar_catalogo_incluir_pendientes_las_devuelve_en_materiales(cliente, db_session, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "codigos_proyecto_path", None)
    expediente, _lote, documento, _linea = _sembrar_catalogo(db_session)
    _agregar_linea_huerfana(db_session, expediente, documento, motivo="ninguna cabecera LOTE N encontrada")

    resp = cliente.get("/catalogo/exportar.xlsx", params={"incluir_pendientes": "true"})

    ruta = tmp_path / "salida.xlsx"
    ruta.write_bytes(resp.content)
    libro = openpyxl.load_workbook(ruta)
    materiales = libro["Materiales"]
    assert materiales.max_row == 3
    descripciones = [c.value for c in materiales["F"]][1:]
    assert "TORNILLO M8" in descripciones


def test_revision_lista_solo_pendientes(cliente, db_session):
    _sembrar_catalogo(db_session, estado=EstadoExpediente.pendiente_revision)

    resp = cliente.get("/revision")

    assert resp.status_code == 200
    datos = resp.json()
    # Paginado (sesión 2026-09-08, mismo criterio que /catalogo): la forma
    # de la respuesta ahora es {total, pagina, tamano_pagina, expedientes}.
    assert datos["total"] == 1
    assert datos["pagina"] == 1
    assert len(datos["expedientes"]) == 1
    assert datos["expedientes"][0]["estado"] == "pendiente_revision"
    assert datos["expedientes"][0]["error"] == "no se pudo determinar la baja del lote"


def test_revision_lista_pagina(cliente, db_session):
    # Sesión de paginación de la cola de revisión (2026-09-08): mismo
    # criterio que /catalogo -- `tamano_pagina` acota cuántos vienen por
    # página, `total` sigue reflejando el recuento real completo.
    for i in range(3):
        db_session.add(
            Expediente(
                codigo_expediente=f"6.24/28510.100{i}",
                estado=EstadoExpediente.pendiente_revision,
                error="motivo de prueba",
            )
        )
    db_session.commit()

    resp = cliente.get("/revision", params={"pagina": 1, "tamano_pagina": 2})
    assert resp.status_code == 200
    datos = resp.json()
    assert datos["total"] == 3
    assert datos["pagina"] == 1
    assert datos["tamano_pagina"] == 2
    assert len(datos["expedientes"]) == 2

    resp2 = cliente.get("/revision", params={"pagina": 2, "tamano_pagina": 2})
    datos2 = resp2.json()
    assert datos2["total"] == 3
    assert len(datos2["expedientes"]) == 1


def test_detalle_revision_incluye_documentos_y_lineas(cliente, db_session):
    expediente, _lote, documento, _linea = _sembrar_catalogo(db_session, estado=EstadoExpediente.pendiente_revision)

    resp = cliente.get(f"/expedientes/{expediente.id}/revision")

    assert resp.status_code == 200
    datos = resp.json()
    assert datos["expediente"]["id"] == expediente.id
    assert len(datos["documentos"]) == 1
    assert datos["documentos"][0]["nombre_archivo"] == "6.24_28510.0008_ANEJO_1.pdf"
    assert len(datos["lineas"]) == 1


def test_confirmar_revision_sin_cuerpo_marca_completado(cliente, db_session):
    expediente, _lote, _doc, _linea = _sembrar_catalogo(db_session, estado=EstadoExpediente.pendiente_revision)

    resp = cliente.post(f"/expedientes/{expediente.id}/revision/confirmar")

    assert resp.status_code == 200
    datos = resp.json()
    assert datos["estado"] == "completado"
    assert datos["error"] is None


def test_confirmar_revision_con_correccion_de_baja_recalcula_precio_adjudicado(cliente, db_session):
    expediente, lote, _doc, linea = _sembrar_catalogo(db_session, estado=EstadoExpediente.pendiente_revision)

    resp = cliente.post(
        f"/expedientes/{expediente.id}/revision/confirmar",
        json={"baja_global": "0.1000"},
    )

    assert resp.status_code == 200
    db_session.refresh(lote)
    db_session.refresh(linea)
    assert lote.baja_lote == Decimal("0.1000")
    assert linea.baja_lote == Decimal("0.1000")
    assert linea.precio_adjudicado == Decimal("21.6000")  # 24,00 * 0,90


def test_confirmar_revision_con_matriz_autorreferenciada_rechaza(cliente, db_session):
    # Cuarta variante real del bug de autorreferencia de matriz
    # (docs/identidad-expediente.md, caso 6.23/28510.0109): la corrección
    # manual desde la cola de revisión es uno de los cinco sitios que
    # escriben `codigo_matriz`, y también tiene que pasar por el único punto
    # de comprobación (`app.extraccion.cruce_codigos.asignar_matriz`).
    expediente, _lote, _doc, _linea = _sembrar_catalogo(db_session, estado=EstadoExpediente.pendiente_revision)

    resp = cliente.post(
        f"/expedientes/{expediente.id}/revision/confirmar",
        json={"codigo_matriz": expediente.codigo_expediente},
    )

    assert resp.status_code == 400
    db_session.refresh(expediente)
    assert expediente.codigo_matriz is None


def test_crear_expediente_con_matriz_autorreferenciada_rechaza(cliente, db_session):
    # Mismo guard que la corrección manual, pero en el alta directa
    # (`POST /expedientes`, routers/expedientes.py) -- otro de los cinco
    # sitios que escribían `codigo_matriz` sin pasar por
    # `asignar_matriz` antes de esta sesión.
    resp = cliente.post(
        "/expedientes",
        json={"codigo_expediente": "6.23/28510.0109", "codigo_matriz": "6.23/28510.0109"},
    )

    assert resp.status_code == 400
    assert db_session.query(Expediente).count() == 0


def test_crear_expediente_con_matriz_valida_la_guarda(cliente, db_session):
    resp = cliente.post(
        "/expedientes",
        json={"codigo_expediente": "6.23/28510.0018", "codigo_matriz": "6.23/28510.0102"},
    )

    assert resp.status_code == 200
    assert resp.json()["codigo_matriz"] == "6.23/28510.0102"


def test_corregir_linea_catalogo_actualiza_y_marca_corregido(cliente, db_session):
    _expediente, _lote, _doc, linea = _sembrar_catalogo(db_session)

    resp = cliente.patch(
        f"/catalogo/lineas/{linea.id}",
        json={"matricula": "697500901", "comentarios": "matrícula corregida a mano"},
    )

    assert resp.status_code == 200
    datos = resp.json()
    assert datos["matricula"] == "697500901"
    assert datos["comentarios"] == "matrícula corregida a mano"
    assert datos["estado_revision"] == "corregido"
    # Un campo no incluido en la corrección no se borra (mismo criterio que
    # guardar_lineas_catalogo, CONTEXTO.md sección 9.9).
    assert datos["descripcion"] == "GUANTE CONTRA RIESGO ELECTRICO"


def test_confirmar_linea_catalogo(cliente, db_session):
    _expediente, _lote, _doc, linea = _sembrar_catalogo(db_session)
    assert linea.estado_revision == EstadoRevisionLinea.sin_revisar

    resp = cliente.post(f"/catalogo/lineas/{linea.id}/confirmar")

    assert resp.status_code == 200
    assert resp.json()["estado_revision"] == "confirmado"


def test_descartar_linea_catalogo_exige_motivo(cliente, db_session):
    _expediente, _lote, _doc, linea = _sembrar_catalogo(db_session)

    resp = cliente.post(f"/catalogo/lineas/{linea.id}/descartar", json={"motivo": ""})

    assert resp.status_code == 422


def test_descartar_linea_catalogo(cliente, db_session):
    # CONTEXTO.md bloque 2: una línea que no es material real, o no se puede
    # determinar, se saca del catálogo con un motivo -- nunca se borra.
    _expediente, _lote, _doc, linea = _sembrar_catalogo(db_session)

    resp = cliente.post(
        f"/catalogo/lineas/{linea.id}/descartar",
        json={"motivo": "es el epígrafe de la tabla, no una línea de material"},
    )

    assert resp.status_code == 200
    datos = resp.json()
    assert datos["estado_revision"] == "descartado"
    assert "es el epígrafe de la tabla" in datos["comentarios"]


def test_marcar_linea_pendiente_exige_nota(cliente, db_session):
    _expediente, _lote, _doc, linea = _sembrar_catalogo(db_session)

    resp = cliente.post(f"/catalogo/lineas/{linea.id}/pendiente", json={"nota": ""})

    assert resp.status_code == 422


def test_marcar_linea_pendiente(cliente, db_session):
    _expediente, _lote, _doc, linea = _sembrar_catalogo(db_session)

    resp = cliente.post(
        f"/catalogo/lineas/{linea.id}/pendiente",
        json={"nota": "consultar con compras si esta matrícula sigue vigente"},
    )

    assert resp.status_code == 200
    datos = resp.json()
    assert datos["estado_revision"] == "pendiente"
    assert "consultar con compras" in datos["comentarios"]


def test_descartar_y_pendiente_acumulan_comentarios_sin_perder_los_anteriores(cliente, db_session):
    _expediente, _lote, _doc, linea = _sembrar_catalogo(db_session)
    cliente.patch(f"/catalogo/lineas/{linea.id}", json={"comentarios": "nota original"})

    resp = cliente.post(f"/catalogo/lineas/{linea.id}/pendiente", json={"nota": "nota nueva"})

    datos = resp.json()
    assert "nota original" in datos["comentarios"]
    assert "nota nueva" in datos["comentarios"]


def test_exportar_catalogo_excluye_lineas_descartadas(cliente, db_session, monkeypatch):
    monkeypatch.setattr(settings, "codigos_proyecto_path", None)
    _expediente, _lote, _doc, linea = _sembrar_catalogo(db_session)
    cliente.post(f"/catalogo/lineas/{linea.id}/descartar", json={"motivo": "no es material real"})

    resp = cliente.get("/catalogo/exportar.xlsx")

    libro = openpyxl.load_workbook(io.BytesIO(resp.content))
    hoja = libro.active
    assert hoja.max_row == 1  # solo la cabecera: la única línea del catálogo estaba descartada


def test_catalogo_sigue_mostrando_lineas_descartadas(cliente, db_session):
    # A diferencia del Excel, la pantalla de catálogo/revisión no oculta las
    # descartadas -- CONTEXTO.md bloque 2 solo pide excluirlas de la entrega.
    _expediente, _lote, _doc, linea = _sembrar_catalogo(db_session)
    cliente.post(f"/catalogo/lineas/{linea.id}/descartar", json={"motivo": "no es material real"})

    resp = cliente.get("/catalogo")

    assert resp.json()["total"] == 1


# Bloque 5, cambios del cliente tras revisar el catálogo: lista de exclusión
# de expedientes (app/exclusion.py) -- "6.24/28510.0008" es el código que
# siembra _sembrar_catalogo.
def test_catalogo_oculta_expediente_de_la_lista_de_exclusion(cliente, db_session, tmp_path, monkeypatch):
    fichero = tmp_path / "exclusion.txt"
    fichero.write_text("# comentario\n6.24/28510.0008\n")
    monkeypatch.setattr(settings, "exclusion_expedientes_path", str(fichero))
    _sembrar_catalogo(db_session)

    resp = cliente.get("/catalogo")

    assert resp.json()["total"] == 0


def test_catalogo_oculta_departamento_completo_de_la_lista_de_exclusion(cliente, db_session, tmp_path, monkeypatch):
    fichero = tmp_path / "exclusion.txt"
    fichero.write_text("28510\n")
    monkeypatch.setattr(settings, "exclusion_expedientes_path", str(fichero))
    _sembrar_catalogo(db_session)

    resp = cliente.get("/catalogo")

    assert resp.json()["total"] == 0


def test_catalogo_no_oculta_expediente_de_otro_departamento(cliente, db_session, tmp_path, monkeypatch):
    fichero = tmp_path / "exclusion.txt"
    fichero.write_text("28520\n")
    monkeypatch.setattr(settings, "exclusion_expedientes_path", str(fichero))
    _sembrar_catalogo(db_session)  # "6.24/28510.0008", departamento 28510

    resp = cliente.get("/catalogo")

    assert resp.json()["total"] == 1


# Bloque 3, sesión 2026-09-09: mecanismo de exclusión por código interno.
def test_catalogo_oculta_codigo_interno_de_la_lista_de_exclusion(cliente, db_session, tmp_path, monkeypatch):
    fichero = tmp_path / "exclusion.txt"
    fichero.write_text("INTERNO:24038\n")
    monkeypatch.setattr(settings, "exclusion_expedientes_path", str(fichero))
    expediente, *_ = _sembrar_catalogo(db_session)
    expediente.codigo_interno = "24038"
    db_session.commit()

    resp = cliente.get("/catalogo")

    assert resp.json()["total"] == 0


def test_catalogo_no_oculta_expediente_sin_codigo_interno_por_lista_de_codigo_interno(
    cliente, db_session, tmp_path, monkeypatch
):
    # El expediente sembrado no tiene codigo_interno (58,9% del corpus real
    # no cruza con el Excel de códigos) -- sin el `or_(... is_(None) ...)` de
    # `_excluir_expedientes_de_la_lista`, "NULL NOT IN (...)" habría ocultado
    # también a este, no solo a los de la lista.
    fichero = tmp_path / "exclusion.txt"
    fichero.write_text("INTERNO:24038\n")
    monkeypatch.setattr(settings, "exclusion_expedientes_path", str(fichero))
    _sembrar_catalogo(db_session)

    resp = cliente.get("/catalogo")

    assert resp.json()["total"] == 1


# Bloque 3, sesión de comparación documento-vs-listado interno: filtro por
# palabras del título del contrato (app/exclusion.py,
# ExclusionPalabrasTitulo) -- "SUMINISTRO DE GUANTES CONTRA RIESGO
# ELECTRICO." es el título que siembra _sembrar_catalogo.
def test_catalogo_oculta_por_palabra_del_titulo(cliente, db_session, tmp_path, monkeypatch):
    fichero = tmp_path / "palabras.txt"
    fichero.write_text("# no es material\nguantes\n")
    monkeypatch.setattr(settings, "exclusion_palabras_titulo_path", str(fichero))
    _sembrar_catalogo(db_session)

    resp = cliente.get("/catalogo")

    assert resp.json()["total"] == 0


def test_catalogo_no_oculta_por_palabra_ajena_al_titulo(cliente, db_session, tmp_path, monkeypatch):
    fichero = tmp_path / "palabras.txt"
    fichero.write_text("arrendamiento\n")
    monkeypatch.setattr(settings, "exclusion_palabras_titulo_path", str(fichero))
    _sembrar_catalogo(db_session)

    resp = cliente.get("/catalogo")

    assert resp.json()["total"] == 1


def test_exportar_catalogo_excluye_por_palabra_del_titulo_pero_no_lo_borra(
    cliente, db_session, tmp_path, monkeypatch
):
    monkeypatch.setattr(settings, "codigos_proyecto_path", None)
    fichero = tmp_path / "palabras.txt"
    fichero.write_text("guantes\n")
    monkeypatch.setattr(settings, "exclusion_palabras_titulo_path", str(fichero))
    expediente, *_ = _sembrar_catalogo(db_session)

    resp = cliente.get("/catalogo/exportar.xlsx")

    libro = openpyxl.load_workbook(io.BytesIO(resp.content))
    hoja = libro.active
    assert hoja.max_row == 1  # solo la cabecera: el único expediente está excluido
    # Se conserva en base de datos, igual que la exclusión por expediente/
    # departamento/código interno de arriba -- la exclusión es solo de la
    # vista (CONTEXTO.md sección 7).
    assert db_session.query(Expediente).filter_by(codigo_expediente=expediente.codigo_expediente).count() == 1


def test_exportar_catalogo_excluye_expediente_de_la_lista_pero_no_lo_borra(
    cliente, db_session, tmp_path, monkeypatch
):
    monkeypatch.setattr(settings, "codigos_proyecto_path", None)
    fichero = tmp_path / "exclusion.txt"
    fichero.write_text("6.24/28510.0008\n")
    monkeypatch.setattr(settings, "exclusion_expedientes_path", str(fichero))
    _sembrar_catalogo(db_session)

    resp = cliente.get("/catalogo/exportar.xlsx")

    libro = openpyxl.load_workbook(io.BytesIO(resp.content))
    hoja = libro.active
    assert hoja.max_row == 1  # solo la cabecera: el único expediente está excluido
    # Se conserva en base de datos (CONTEXTO.md, encargo del bloque 5: "se
    # conservan en la base de datos") -- la exclusión es solo de la vista.
    assert db_session.query(Expediente).filter_by(codigo_expediente="6.24/28510.0008").count() == 1


def test_descargar_documento_sirve_los_bytes_reales(cliente, db_session, tmp_path, monkeypatch):
    from app.routers import documentos as documentos_router

    _expediente, _lote, documento, _linea = _sembrar_catalogo(db_session)
    storage = LocalDiskStorage(str(tmp_path))
    storage.guardar(documento.ruta_almacenamiento, b"%PDF-1.4 contenido de prueba")
    monkeypatch.setattr(documentos_router, "_storage", storage)

    resp = cliente.get(f"/documentos/{documento.id}/archivo")

    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.content == b"%PDF-1.4 contenido de prueba"


def test_descargar_documento_inexistente_da_404(cliente, db_session):
    resp = cliente.get("/documentos/999/archivo")
    assert resp.status_code == 404
